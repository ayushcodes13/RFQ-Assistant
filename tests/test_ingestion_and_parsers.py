"""Unit and integration tests for document validation, storage, and deterministic parsers."""

from pathlib import Path
import shutil
import pytest

from rfq_assistant.ingestion import (
    DocumentParserError,
    DocumentStorage,
    UploadValidationError,
    compute_sha256,
    get_parser,
    sanitize_filename,
    validate_file_upload,
)
from rfq_assistant.schemas.enums import SourceType

ROOT = Path(__file__).resolve().parent.parent
SAMPLES_DIR = ROOT / "fixtures" / "rfq_samples"


class TestUploadValidationAndStorage:
    """Verifies security controls, file limits, sanitization, and persistence."""

    def test_sanitize_filename_prevents_traversal(self):
        assert sanitize_filename("../../../etc/passwd") == "passwd"
        assert sanitize_filename("..\\..\\windows\\system32\\cmd.exe") == "cmd.exe"
        assert sanitize_filename(".hidden_file.xlsx") == "upload_.hidden_file.xlsx"

    def test_reject_empty_file(self):
        with pytest.raises(UploadValidationError, match="empty"):
            validate_file_upload("empty.xlsx", b"")

    def test_reject_oversized_file(self):
        oversized = b"A" * (10 * 1024 * 1024 + 1)
        with pytest.raises(UploadValidationError, match="exceeds maximum permitted limit"):
            validate_file_upload("big.csv", oversized)

    def test_reject_macro_spreadsheet(self):
        with pytest.raises(UploadValidationError, match="Macro-enabled spreadsheets"):
            validate_file_upload("bom_with_vba.xlsm", b"DUMMY_CONTENT")

    def test_reject_unsupported_extension(self):
        with pytest.raises(UploadValidationError, match="Unsupported file format"):
            validate_file_upload("script.py", b"print('hello')")

    def test_valid_csv_upload(self):
        content = b"Part Number,Manufacturer,Qty\nLC1D09BD,Schneider Electric,10\n"
        name, src_type, digest = validate_file_upload("parts.csv", content)
        assert name == "parts.csv"
        assert src_type == SourceType.CSV
        assert digest == compute_sha256(content)

    def test_document_storage_save_and_retrieve(self, tmp_path):
        storage = DocumentStorage(base_dir=tmp_path)
        content = b"Sample RFQ Data"
        saved_path = storage.save_file("case-123", "rfq.txt", content)
        assert saved_path.exists()
        retrieved = storage.get_file("case-123", "rfq.txt")
        assert retrieved == content
        storage.delete_case("case-123")
        assert not saved_path.exists()

    def test_document_storage_rejects_sibling_prefix_traversal(self, tmp_path):
        storage = DocumentStorage(base_dir=tmp_path)
        sibling = storage.uploads_dir / "case-123-escape"
        sibling.mkdir()

        with pytest.raises(ValueError, match="Path traversal blocked"):
            storage.save_file("case-123", "../case-123-escape/payload.txt", b"synthetic")

        assert not (sibling / "payload.txt").exists()


class TestSpreadsheetParser:
    """Verifies deterministic parsing of XLSX and CSV files with multi-sheet and edge case handling."""

    def test_parse_clean_20line_bom(self):
        file_path = SAMPLES_DIR / "industrial_control_bom_20lines.xlsx"
        content = file_path.read_bytes()
        parser = get_parser(SourceType.XLSX)
        result = parser.parse(content, "industrial_control_bom_20lines.xlsx")

        assert result.total_lines_detected == 20
        assert len(result.lines) == 20

        # Check line 1
        line1 = result.lines[0]
        assert line1.line_number == 1
        assert line1.requested_mpn == "LC1D09BD"
        assert line1.requested_manufacturer == "Schneider Electric"
        assert line1.quantity == 10.0
        assert line1.unit == "EA"
        assert line1.source_locator.sheet_name == "BOM Items"
        assert line1.source_locator.row_number == 5

        # Check line 18: missing quantity must NOT be dropped, but must carry a warning
        line18 = result.lines[17]
        assert line18.line_number == 18
        assert line18.requested_mpn == "UT 2.5"
        assert line18.requested_manufacturer == "Phoenix Contact"
        assert line18.quantity is None
        assert any(w.code == "MISSING_QUANTITY" for w in line18.extraction_warnings)

        # Check line 19: non-catalog item
        line19 = result.lines[18]
        assert line19.requested_mpn == "NONEXISTENT-PART-999"
        assert line19.quantity == 2.0

        # Check line 20: suffix part
        line20 = result.lines[19]
        assert line20.requested_mpn == "ABC-12A"
        assert line20.requested_manufacturer == "Acme Industrial"

    def test_parse_multisheet_project_bom(self):
        file_path = SAMPLES_DIR / "multisheet_project_bom.xlsx"
        content = file_path.read_bytes()
        parser = get_parser(SourceType.XLSX)
        result = parser.parse(content, "multisheet_project_bom.xlsx")

        assert result.total_lines_detected == 5
        sheet_names = {line.source_locator.sheet_name for line in result.lines}
        assert sheet_names == {"Control Panel A", "Field Instrumentation"}

        # Verify items from sheet 1
        panel_a_lines = [l for l in result.lines if l.source_locator.sheet_name == "Control Panel A"]
        assert len(panel_a_lines) == 3

        # Verify items from sheet 2
        field_lines = [l for l in result.lines if l.source_locator.sheet_name == "Field Instrumentation"]
        assert len(field_lines) == 2

    def test_xlsx_formula_quantity_is_preserved_but_not_calculated(self):
        import io
        import openpyxl
        from rfq_assistant.ingestion.parsers.spreadsheet_parser import SpreadsheetParser

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["MPN", "Manufacturer", "Quantity"])
        sheet.append(["LC1D09BD", "Schneider Electric", "=5+5"])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        parsed = SpreadsheetParser().parse(output.getvalue(), "formula_quantity.xlsx")
        assert len(parsed.lines) == 1
        line = parsed.lines[0]
        assert line.quantity is None
        assert "=5+5" in line.raw_text
        assert any(w.code == "FORMULA_VALUE_NOT_EVALUATED" for w in line.extraction_warnings)
        assert any(w.code == "MISSING_QUANTITY" for w in line.extraction_warnings)
        assert len(parsed.source_coverage) == 1
        sheet_coverage = parsed.source_coverage[0]
        assert sheet_coverage.unit_id == "sheet:Sheet"
        assert sheet_coverage.nonempty_regions == 2
        assert sheet_coverage.structured_regions == 1
        assert sheet_coverage.filtered_regions == 1
        assert sheet_coverage.header_row_number == 1
        assert sheet_coverage.formula_backed_fields == 1

    def test_xlsx_formula_cached_value_is_shown_but_never_used_as_quantity(self):
        import io
        import zipfile
        from xml.etree import ElementTree

        import openpyxl
        from rfq_assistant.ingestion.parsers.spreadsheet_parser import SpreadsheetParser

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["MPN", "Manufacturer", "Quantity"])
        sheet.append(["LC1D09BD", "Schneider Electric", "=5+5"])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        # openpyxl writes formulas without cached results. Insert a synthetic
        # workbook cache value into the generated fixture's XML for this test.
        cached_fixture = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(output.getvalue())) as source_zip:
            with zipfile.ZipFile(cached_fixture, "w", zipfile.ZIP_DEFLATED) as target_zip:
                for entry in source_zip.infolist():
                    data = source_zip.read(entry.filename)
                    if entry.filename == "xl/worksheets/sheet1.xml":
                        root = ElementTree.fromstring(data)
                        ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                        cell = root.find(".//x:c[@r='C2']", ns)
                        assert cell is not None
                        value = cell.find("x:v", ns)
                        assert value is not None
                        value.text = "10"
                        data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
                    target_zip.writestr(entry, data)

        parsed = SpreadsheetParser().parse(cached_fixture.getvalue(), "formula_quantity.xlsx")
        line = parsed.lines[0]
        warning = next(w for w in line.extraction_warnings if w.code == "FORMULA_VALUE_NOT_EVALUATED")
        assert line.quantity is None
        assert "cached value" not in line.raw_text
        assert "quantity=10" in warning.message

    def test_xlsx_zero_padding_format_preserves_numeric_mpn_leading_zeroes(self):
        import io
        import openpyxl
        from rfq_assistant.ingestion.parsers.spreadsheet_parser import SpreadsheetParser

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["MPN", "Manufacturer", "Quantity"])
        sheet.append([2900364, "Phoenix Contact", 2])
        sheet["A2"].number_format = "00000000"
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        parsed = SpreadsheetParser().parse(output.getvalue(), "numeric_mpn.xlsx")
        assert parsed.lines[0].requested_mpn == "02900364"
        assert "02900364" in parsed.lines[0].raw_text

    def test_xlsx_numeric_mpn_over_excel_precision_limit_is_flagged(self):
        import io
        import openpyxl
        from rfq_assistant.ingestion.parsers.spreadsheet_parser import SpreadsheetParser

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["MPN", "Manufacturer", "Quantity"])
        sheet.append([1234567890123456, "Synthetic Manufacturer", 1])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        parsed = SpreadsheetParser().parse(output.getvalue(), "long_numeric_mpn.xlsx")
        assert any(
            warning.code == "NUMERIC_MPN_EXCEL_PRECISION_RISK"
            for warning in parsed.lines[0].extraction_warnings
        )

    def test_xlsx_hidden_rows_and_mapped_hidden_columns_are_manifested_and_flagged(self):
        import io
        import openpyxl
        from rfq_assistant.ingestion.parsers.spreadsheet_parser import SpreadsheetParser

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["MPN", "Manufacturer", "Quantity"])
        sheet.append(["LC1D09BD", "Schneider Electric", 2])
        sheet.row_dimensions[2].hidden = True
        sheet.column_dimensions["A"].hidden = True
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        parsed = SpreadsheetParser().parse(output.getvalue(), "hidden_layout.xlsx")
        warning_codes = {warning.code for warning in parsed.lines[0].extraction_warnings}
        coverage = next(unit for unit in parsed.source_coverage if unit.unit_kind == "workbook_sheet")

        assert "HIDDEN_ROW_REVIEW" in warning_codes
        assert "HIDDEN_COLUMN_VALUE_REVIEW" in warning_codes
        assert coverage.hidden_row_count == 1
        assert coverage.hidden_column_count == 1


class TestPDFParser:
    """Verifies native PDF parsing and safe detection of corrupt/scanned documents."""

    def test_parse_native_pdf(self):
        file_path = SAMPLES_DIR / "formal_rfq_native.pdf"
        content = file_path.read_bytes()
        parser = get_parser(SourceType.PDF_NATIVE)
        result = parser.parse(content, "formal_rfq_native.pdf")

        assert result.total_lines_detected >= 7
        mpns = [l.requested_mpn for l in result.lines]
        assert "LC1D09BD" in mpns
        assert "6ES7214-1AG40-0XB0" in mpns
        # Digit-only component IDs must be kept; trailing electrical
        # specifications such as 24VDC are not the part number.
        phoenix_line = next(line for line in result.lines if line.requested_manufacturer == "Phoenix Contact")
        assert phoenix_line.requested_mpn == "2900364"
        assert all(l.source_locator.page_number == 1 for l in result.lines)
        coverage = result.source_coverage
        assert [unit.unit_id for unit in coverage] == ["page:1"]
        assert coverage[0].nonempty_regions == (
            coverage[0].structured_regions
            + coverage[0].review_regions
            + coverage[0].filtered_regions
        )
        assert len(coverage[0].content_sha256) == 64

    def test_corrupt_pdf_fails_safely(self):
        file_path = SAMPLES_DIR / "corrupt_file.pdf"
        content = file_path.read_bytes()
        parser = get_parser(SourceType.PDF_NATIVE)
        with pytest.raises(DocumentParserError, match="Corrupted or invalid PDF"):
            parser.parse(content, "corrupt_file.pdf")

    def test_scanned_pdf_rejection_when_ocr_disabled(self):
        # A valid empty PDF without text layer simulates a scanned document
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        empty_pdf_bytes = bytes(pdf.output())

        parser = get_parser(SourceType.PDF_NATIVE, allow_ocr=False)
        with pytest.raises(DocumentParserError, match="Scanned or image-only PDF detected"):
            parser.parse(empty_pdf_bytes, "scanned.pdf")

    def test_unparsed_numbered_product_row_is_retained_for_fallback_or_manual_review(self):
        from fpdf import FPDF

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 10, "1 safety disconnect switch")
        result = get_parser(SourceType.PDF_NATIVE, allow_ocr=False).parse(
            bytes(pdf.output()), "unusual_native_rfq.pdf"
        )

        assert result.lines == []
        assert len(result.candidate_segments) == 1
        candidate = result.candidate_segments[0]
        assert candidate.locator.page_number == 1
        assert "disconnect switch" in candidate.text

    def test_unfamiliar_nonempty_pdf_text_is_retained_even_without_product_keywords(self):
        from fpdf import FPDF

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 10, "Need a custom assembly for the remote panel")
        result = get_parser(SourceType.PDF_NATIVE, allow_ocr=False).parse(
            bytes(pdf.output()), "unfamiliar_rfq.pdf"
        )

        assert len(result.candidate_segments) == 1
        assert result.candidate_segments[0].locator.page_number == 1
        assert "custom assembly" in result.candidate_segments[0].text

    def test_unfamiliar_text_region_is_retained_without_known_item_keywords(self):
        from rfq_assistant.ingestion.parsers.text_parser import TextParser

        result = TextParser().parse(
            b"Please quote a custom assembly for the remote panel, quantity 7",
            "unfamiliar_rfq.txt",
        )

        assert result.lines == []
        assert len(result.candidate_segments) == 1
        assert result.candidate_segments[0].locator.row_number == 1

    def test_text_coverage_accounts_for_extraction_review_and_metadata(self):
        from rfq_assistant.ingestion.parsers.text_parser import TextParser

        content = (
            "Subject: Synthetic request\n"
            "1. 10 pcs Schneider Electric LC1D09BD contactor\n"
            "Could you also include the unusual assembly from last time?\n"
        )
        result = TextParser().parse(content.encode(), "coverage.txt")
        assert len(result.source_coverage) == 1
        unit = result.source_coverage[0]
        assert unit.unit_id == "document"
        assert unit.nonempty_regions == 3
        assert unit.structured_regions == 1
        assert unit.review_regions == 1
        assert unit.filtered_regions == 1
        assert unit.structured_regions + unit.review_regions + unit.filtered_regions == unit.nonempty_regions
        assert unit.locator.row_number is None
        assert unit.locator.text_snippet is None
        assert len(unit.content_sha256) == 64
        repeated = TextParser().parse(content.encode(), "coverage.txt")
        assert repeated.source_coverage == result.source_coverage
        changed = TextParser().parse(content.replace("10 pcs", "11 pcs").encode(), "coverage.txt")
        assert changed.source_coverage[0].content_sha256 != unit.content_sha256

    def test_xlsx_manifest_includes_empty_and_hidden_sheets(self):
        import io
        import openpyxl
        from rfq_assistant.ingestion.parsers.spreadsheet_parser import SpreadsheetParser

        workbook = openpyxl.Workbook()
        workbook.active.title = "Empty"
        hidden = workbook.create_sheet("Hidden RFQ")
        hidden.sheet_state = "hidden"
        hidden.append(["Synthetic quotation request"])
        hidden.append(["MPN", "Quantity", "Description"])
        hidden.append(["LC1D09BD", 2, "Contactor"])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        parsed = SpreadsheetParser().parse(output.getvalue(), "hidden_sheet.xlsx")
        coverage = {unit.locator.sheet_name: unit for unit in parsed.source_coverage}
        assert coverage["Empty"].is_empty is True
        assert coverage["Empty"].sheet_visibility == "visible"
        assert coverage["Hidden RFQ"].is_empty is False
        assert coverage["Hidden RFQ"].sheet_visibility == "hidden"
        assert coverage["Hidden RFQ"].header_row_number == 2

    def test_candidate_region_limit_fails_explicitly(self):
        from rfq_assistant.ingestion.parsers.text_parser import TextParser

        text = "\n".join(f"Unrecognized request detail {idx}" for idx in range(501))
        with pytest.raises(DocumentParserError, match="more than 500 unresolved"):
            TextParser().parse(text.encode(), "many_regions.txt")

    def test_sparse_image_page_uses_bounded_local_ocr(self, monkeypatch):
        """Only image-heavy pages are OCRed and every OCR-derived row stays reviewable."""
        import io
        import fitz
        from PIL import Image
        from rfq_assistant.ingestion.parsers import pdf_parser

        image_buffer = io.BytesIO()
        Image.new("RGB", (600, 800), "white").save(image_buffer, format="PNG")
        pdf = fitz.open()
        page = pdf.new_page(width=595, height=842)
        page.insert_image(page.rect, stream=image_buffer.getvalue())
        content = pdf.tobytes()
        pdf.close()

        monkeypatch.setattr(
            pdf_parser,
            "recognize_page",
            lambda _content, _page_number: "1 Siemens 6ES7214-1AG40-0XB0 PLC 2 EA",
        )
        result = get_parser(SourceType.PDF_NATIVE, allow_ocr=True).parse(content, "scan.pdf")

        assert len(result.lines) == 1
        line = result.lines[0]
        assert line.requested_mpn == "6ES7214-1AG40-0XB0"
        assert line.quantity == 2.0
        assert line.source_locator.page_number == 1
        assert any(w.code == "OCR_TEXT_REVIEW" for w in line.extraction_warnings)
        assert any(w.code == "LOCAL_OCR_USED" for w in result.warnings)
        page_evidence = next(segment for segment in result.evidence_segments if segment.locator.page_number == 1)
        assert "Siemens" in page_evidence.text
        assert any(w.code == "OCR_TEXT_REVIEW" for w in page_evidence.warnings)

    def test_page_triage_failure_sends_every_page_to_ocr(self, monkeypatch):
        """An unknown page classification must not silently trust native text alone."""
        from fpdf import FPDF
        from rfq_assistant.ingestion.parsers import pdf_parser

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(0, 10, text="1 Siemens 6ES7214-1AG40-0XB0 PLC 2 EA")
        content = bytes(pdf.output())
        calls = []

        def fail_triage(_page, _text):
            raise RuntimeError("synthetic triage failure")

        def fake_ocr(_content, page_number):
            calls.append(page_number)
            return "1 Siemens 6ES7214-1AG40-0XB0 PLC 2 EA"

        monkeypatch.setattr(pdf_parser, "page_needs_ocr", fail_triage)
        monkeypatch.setattr(pdf_parser, "recognize_page", fake_ocr)
        result = pdf_parser.PDFParser(allow_ocr=True).parse(content, "triage-fallback.pdf")

        assert calls == [1]
        assert result.lines[0].requested_mpn == "6ES7214-1AG40-0XB0"
        assert any(w.code == "PDF_PAGE_TRIAGE_UNAVAILABLE" for w in result.warnings)
        assert any(w.code == "LOCAL_OCR_USED" for w in result.warnings)
        assert any(w.code == "OCR_TEXT_REVIEW" for w in result.lines[0].extraction_warnings)

    @pytest.mark.skipif(
        shutil.which("tesseract") is None,
        reason="local OCR integration requires the Tesseract binary",
    )
    @pytest.mark.parametrize("rotation", [90, 270])
    def test_rotated_scan_is_ocrd_in_reading_orientation_with_timing(self, rotation):
        """A rotated image-only page keeps readable OCR and rotation provenance."""
        import io
        import fitz
        from PIL import Image
        from rfq_assistant.matching import ExactMatchEngine
        from rfq_assistant.matching.catalog import CatalogRepository
        from rfq_assistant.ingestion.parsers.pdf_parser import PDFParser
        from rfq_assistant.schemas.enums import MatchStatus
        from rfq_assistant.schemas.models import CatalogItem

        source = fitz.open()
        source_page = source.new_page(width=1400, height=400)
        source_page.insert_text(
            (40, 210), "1 Siemens 6ES7214-1AG40-0XB0 PLC 2 EA",
            fontsize=56, fontname="helv",
        )
        upright_png = source_page.get_pixmap(
            matrix=fitz.Matrix(2, 2), alpha=False
        ).tobytes("png")
        source.close()

        sideways = Image.open(io.BytesIO(upright_png)).rotate(
            90 if rotation == 90 else 270, expand=True
        )
        image_buffer = io.BytesIO()
        sideways.save(image_buffer, format="PNG")
        pdf = fitz.open()
        page = pdf.new_page(width=400, height=1400)
        page.insert_image(page.rect, stream=image_buffer.getvalue())
        page.set_rotation(rotation)
        content = pdf.tobytes()
        pdf.close()

        parsed = PDFParser(allow_ocr=True).parse(content, "rotated_scan.pdf")
        page_coverage = next(
            unit for unit in parsed.source_coverage
            if unit.unit_kind == "pdf_page"
        )

        assert page_coverage.page_rotation_degrees == rotation
        assert page_coverage.ocr_status == "succeeded"
        assert page_coverage.ocr_duration_ms is not None
        assert page_coverage.ocr_duration_ms > 0
        assert page_coverage.locator.page_bbox is None
        assert len(parsed.lines) == 1
        assert parsed.lines[0].requested_manufacturer == "Siemens"
        assert parsed.lines[0].quantity == 2
        bbox = parsed.lines[0].source_locator.page_bbox
        assert bbox is not None
        assert parsed.lines[0].source_locator.row_number == 1
        assert 0 <= bbox[0] < bbox[2] <= 1
        assert 0 <= bbox[1] < bbox[3] <= 1
        assert parsed.lines[0].source_locator.ocr_confidence is not None
        catalog = CatalogRepository()
        catalog.add_item(CatalogItem(
            internal_sku="SIEMENS-PLC-1",
            manufacturer="Siemens",
            manufacturer_part_number="6ES7214-1AG40-0XB0",
            description="Synthetic PLC",
        ))
        match = ExactMatchEngine(catalog).match_line(
            parsed.lines[0].requested_mpn,
            parsed.lines[0].requested_manufacturer,
        )
        assert match.match_status == MatchStatus.NO_MATCH
        assert any(
            warning.code == "OCR_TEXT_REVIEW"
            for warning in parsed.lines[0].extraction_warnings
        )

    def test_image_heavy_page_is_blocked_when_ocr_disabled(self):
        import io
        import fitz
        from PIL import Image

        image_buffer = io.BytesIO()
        Image.new("RGB", (600, 800), "white").save(image_buffer, format="PNG")
        pdf = fitz.open()
        page = pdf.new_page(width=595, height=842)
        page.insert_image(page.rect, stream=image_buffer.getvalue())
        content = pdf.tobytes()
        pdf.close()

        with pytest.raises(DocumentParserError, match="image|OCR") as error:
            get_parser(SourceType.PDF_NATIVE, allow_ocr=False).parse(content, "scan.pdf")
        page = error.value.source_coverage[0]
        assert page.ocr_status == "disabled"
        assert page.processing_status == "needs_review"
        assert page.is_empty is True


class TestTextAndEmailParser:
    """Verifies deterministic parsing of messy email bodies, vague references, and security inputs."""

    def test_parse_messy_email_quote_request(self):
        file_path = SAMPLES_DIR / "messy_email_quote_request.txt"
        content = file_path.read_bytes()
        parser = get_parser(SourceType.EMAIL_TEXT)
        result = parser.parse(content, "messy_email_quote_request.txt")

        assert result.total_lines_detected == 6
        mpns = [l.requested_mpn for l in result.lines]
        assert "LC1D09BD" in mpns
        assert "6ES7223-1BL32-0XB0" in mpns
        assert "XB4BA31" in mpns
        assert "MY4N-D2-DC24" in mpns
        assert "FAZ-C10/1" in mpns

        # Item 5 is "sensor from last time" -> must be flagged with UNKNOWN_PART_REFERENCE
        vague_line = [l for l in result.lines if l.requested_mpn is None][0]
        assert any(w.code == "UNKNOWN_PART_REFERENCE" for w in vague_line.extraction_warnings)

    def test_adversarial_prompt_injection_safety(self):
        file_path = SAMPLES_DIR / "adversarial_prompt_injection.txt"
        content = file_path.read_bytes()
        parser = get_parser(SourceType.EMAIL_TEXT)
        result = parser.parse(content, "adversarial_prompt_injection.txt")

        # Parsing completes safely as raw strings without executing any injection instruction
        assert result.total_lines_detected >= 3
        # Ensure no arbitrary prices or states were modified
        for line in result.lines:
            assert line.quantity is None or line.quantity > 0
