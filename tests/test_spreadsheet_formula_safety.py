"""Synthetic adversarial tests against spreadsheet-formula execution in RFQ exports."""

import csv
import io

import openpyxl
import pytest

from rfq_assistant.cases.manager import CaseManager
from rfq_assistant.exports.generator import ExportGenerator, ExportValidationError, _safe_csv_cell
from rfq_assistant.ingestion.storage import DocumentStorage
from rfq_assistant.matching.catalog import CatalogRepository
from rfq_assistant.schemas.enums import ReviewDecision
from rfq_assistant.schemas.models import CatalogItem, ExportRequest, RFQLineItemUpdate


@pytest.mark.parametrize("input_value", ["=1+1", "+1+2", "-CMD", "@SUM(1,2)", " \t=1+1"])
def test_dangerous_csv_text_gets_literal_prefix(input_value):
    assert _safe_csv_cell(input_value) == "'" + input_value


def test_clean_fields_and_numeric_values_keep_original_types():
    assert _safe_csv_cell("LC1D09BD") == "LC1D09BD"
    assert _safe_csv_cell(4) == 4
    assert _safe_csv_cell(-5.0) == -5.0


@pytest.fixture
def quote_case(tmp_path):
    catalog = CatalogRepository()
    catalog.add_item(CatalogItem(
        internal_sku="SKU-SYN",
        manufacturer="Schneider Electric",
        manufacturer_part_number="LC1D09BD",
        description="=1+1",  # Untrusted catalog description
        list_price=17.0,
    ))
    manager = CaseManager(catalog=catalog, storage=DocumentStorage(tmp_path / "sources"))
    exporter = ExportGenerator(catalog=catalog, base_export_dir=tmp_path / "exports")
    case = manager.create_case_from_text(
        "10 pcs Schneider Electric LC1D09BD",
        customer_label="=1+1",
    )
    line = manager.get_lines(case.case_id)[0]
    manager.update_line(case.case_id, line.line_id, RFQLineItemUpdate(expected_version=line.version,
        notes="=HYPERLINK(\"https://example.invalid\",\"test\")",
        review_decision=ReviewDecision.APPROVED,
    ))
    # Simulate untrusted RFQ field surviving parsing into a source text field.
    line.requested_manufacturer = "+cmd"
    return exporter, case, manager.get_lines(case.case_id), tmp_path


def test_csv_export_escapes_all_user_supplied_formula_fields(quote_case):
    exporter, case, lines, tmp_path = quote_case
    manifest = exporter.generate_export(
        case, lines, ExportRequest(case_id=case.case_id, format="csv")
    )
    path = tmp_path / "exports" / manifest.export_id / manifest.file_name
    with path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    row = next(row for row in rows if row and row[0] == "1")
    assert row[2] == "'+cmd"
    assert row[7] == "'=1+1"
    assert row[12].startswith("'=HYPERLINK(")


def test_xlsx_export_keeps_suspicious_values_as_literal_strings(quote_case):
    exporter, case, lines, tmp_path = quote_case
    manifest = exporter.generate_export(
        case, lines, ExportRequest(case_id=case.case_id, format="xlsx")
    )
    path = tmp_path / "exports" / manifest.export_id / manifest.file_name
    ws = openpyxl.load_workbook(path, data_only=False).active
    for address in ("E2", "C6", "H6", "M6"):
        assert ws[address].data_type == "s", address
        assert ws[address].value.startswith(("=", "+")), address


def test_zero_line_export_is_blocked_even_if_partial(tmp_path):
    catalog = CatalogRepository()
    generator = ExportGenerator(catalog=catalog, base_export_dir=tmp_path)
    manager = CaseManager(catalog=catalog, storage=DocumentStorage(tmp_path / "docs"))
    case = manager.create_case_from_text("10 pcs Schneider Electric LC1D09BD")
    with pytest.raises(ExportValidationError, match="no extracted line items"):
        generator.generate_export(
            case, [], ExportRequest(case_id=case.case_id, allow_partial=True)
        )
