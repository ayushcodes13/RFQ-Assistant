"""Offline tests for source-grounded optional LLM extraction."""

from __future__ import annotations

import json
from typing import List

import pytest

from rfq_assistant.cases.manager import CaseManager
from rfq_assistant.ingestion.ai_extraction import (
    EvidenceSegment,
    GroundedLLMExtractor,
    OpenAICompatibleClient,
)
from rfq_assistant.ingestion.parsers.base import ParsedDocumentResult
from rfq_assistant.ingestion.storage import DocumentStorage
from rfq_assistant.ingestion.source_segments import SourceSegmentLimitError, source_segments
from rfq_assistant.matching.catalog import CatalogRepository
from rfq_assistant.schemas.models import CatalogItem, SourceLocator


class ScriptedClient:
    def __init__(self, responses: List[str]):
        self.responses = list(responses)
        self.calls = []
        self.model = "gpt-6-luna"

    def complete(self, *, system: str, user: str) -> str:
        self.calls.append((system, user))
        return self.responses.pop(0)


def response(*, quantity=10, manufacturer="Fabrikam Controls", description="relay", quote="Fabrikam Controls ABC-123 10 EA relay", segment_id="row-1"):
    return json.dumps({"lines": [{
        "line_number": 1,
        "requested_manufacturer": manufacturer,
        "requested_mpn": "ABC-123",
        "description": description,
        "quantity": quantity,
        "unit": "EA",
        "segment_id": segment_id,
        "source_quote": quote,
    }]})


def review_response(verdict="consistent"):
    return json.dumps({"lines": [{"line_number": 1, "verdict": verdict}]})


def segment():
    text = "Fabrikam Controls ABC-123 10 EA relay"
    return EvidenceSegment(
        segment_id="row-1",
        locator=SourceLocator(file_name="rfq.txt", row_number=1),
        text=text,
    )


def test_valid_json_is_typed_and_grounded():
    client = ScriptedClient([response(), review_response()])
    result = GroundedLLMExtractor(client).extract([segment()])

    assert len(result.lines) == 1
    line = result.lines[0]
    assert line.requested_mpn == "ABC-123"
    assert line.requested_manufacturer == "Fabrikam Controls"
    assert line.quantity == 10
    assert line.source_locator.row_number == 1
    assert line.field_origins["requested_mpn"] == "LUNA"
    assert line.field_origins["quantity"] == "LUNA"
    assert any(w.code == "LLM_PROPOSED_EXTRACTION" for w in line.extraction_warnings)
    assert len(client.calls) == 2
    assert "Review extracted RFQ line proposals" in client.calls[1][0]
    assert result.llm_model == "gpt-6-luna"
    assert len(result.llm_prompt_sha256) == 64
    assert result.llm_call_count == 2
    assert result.llm_repair_call_count == 0
    assert result.llm_review_call_count == 1


def test_model_cache_requires_explicit_immutable_deployment_version():
    unpinned = GroundedLLMExtractor(ScriptedClient([]))
    assert unpinned.cache_profile() is None

    pinned_client = ScriptedClient([])
    pinned_client.deployment_version = "deployment-revision-2026-10"
    profile = GroundedLLMExtractor(pinned_client).cache_profile()
    assert profile["deployment_version"] == "deployment-revision-2026-10"
    assert profile["schema_version"] == "rfq-grounded-extraction-v1"
    assert profile["validator_version"] == "json-numeric-verbatim-locator-v1"


def test_token_usage_is_aggregated_per_extraction_not_shared_between_calls():
    class UsageClient(ScriptedClient):
        last_usage = None

        def complete(self, *, system: str, user: str) -> str:
            result = super().complete(system=system, user=user)
            self.last_usage = {"input_tokens": 100, "output_tokens": 10}
            return result

    result = GroundedLLMExtractor(UsageClient([response(), review_response()])).extract([segment()])

    assert result.llm_input_tokens == 200
    assert result.llm_output_tokens == 20


def test_ungrounded_quantity_is_rejected_after_single_repair_attempt():
    bad = response(quantity=999)
    client = ScriptedClient([bad, bad])
    result = GroundedLLMExtractor(client).extract([segment()])

    assert result.lines == []
    assert result.warnings[0].code == "LLM_EXTRACTION_BLOCKED"
    assert result.warnings[0].severity == "error"
    assert len(client.calls) == 2
    assert result.llm_call_count == 2
    assert result.llm_repair_call_count == 1


def test_part_number_digits_cannot_ground_a_quantity_without_unit_or_quantity_label():
    quote = "ABC-123"
    bad = json.dumps({"lines": [{
        "line_number": 1,
        "requested_manufacturer": None,
        "requested_mpn": "ABC-123",
        "description": None,
        "quantity": 123,
        "unit": None,
        "segment_id": "row-1",
        "source_quote": quote,
    }]})
    client = ScriptedClient([bad, bad])
    result = GroundedLLMExtractor(client).extract([segment().model_copy(update={"text": quote})])

    assert result.lines == []
    assert result.warnings[0].code == "LLM_EXTRACTION_BLOCKED"


def test_string_quantity_is_rejected_even_if_it_looks_numeric():
    bad = json.loads(response())
    bad["lines"][0]["quantity"] = "10"
    client = ScriptedClient([json.dumps(bad), response(), review_response()])
    result = GroundedLLMExtractor(client).extract([segment()])

    assert result.lines[0].quantity == 10
    assert len(client.calls) == 3


def test_quantity_that_overflows_float_storage_is_rejected():
    huge = "9" * 400
    raw = (
        '{"lines":[{"line_number":1,"requested_manufacturer":"Fabrikam Controls",'
        '"requested_mpn":"ABC-123","description":"relay","quantity":' + huge +
        ',"unit":"EA","segment_id":"row-1","source_quote":"Fabrikam Controls ABC-123 '
        + huge + ' EA relay"}]}'
    )
    client = ScriptedClient([raw, raw])
    result = GroundedLLMExtractor(client).extract([
        segment().model_copy(update={"text": f"Fabrikam Controls ABC-123 {huge} EA relay"})
    ])

    assert result.lines == []
    assert result.warnings[0].code == "LLM_EXTRACTION_BLOCKED"
    assert "represented safely" in result.warnings[0].message


def test_quantity_precision_loss_in_float_storage_is_rejected():
    exact_integer = 9007199254740993
    raw = (
        '{"lines":[{"line_number":1,"requested_manufacturer":"Fabrikam Controls",'
        '"requested_mpn":"ABC-123","description":"relay","quantity":' + str(exact_integer) +
        ',"unit":"EA","segment_id":"row-1","source_quote":"Fabrikam Controls ABC-123 '
        + str(exact_integer) + ' EA relay"}]}'
    )
    client = ScriptedClient([raw, raw])
    result = GroundedLLMExtractor(client).extract([
        segment().model_copy(update={"text": f"Fabrikam Controls ABC-123 {exact_integer} EA relay"})
    ])

    assert result.lines == []
    assert result.warnings[0].code == "LLM_EXTRACTION_BLOCKED"
    assert "represented safely" in result.warnings[0].message


def test_model_cannot_invent_noncontiguous_line_number_ordinals():
    bad = json.loads(response())
    bad["lines"][0]["line_number"] = 99
    client = ScriptedClient([json.dumps(bad), json.dumps(bad)])
    result = GroundedLLMExtractor(client).extract([segment()])

    assert result.lines == []
    assert result.warnings[0].code == "LLM_EXTRACTION_BLOCKED"
    assert len(client.calls) == 2


def test_llm_fills_missing_manufacturer_but_keeps_human_review_warning(tmp_path):
    text = "Fabrikam Controls ABC-123 10 EA relay"
    client = ScriptedClient([response(quote=text, segment_id="draft-line-1"), review_response()])
    manager = CaseManager(
        catalog=CatalogRepository(),
        storage=DocumentStorage(base_dir=tmp_path),
        llm_extractor=GroundedLLMExtractor(client),
    )

    case = manager.create_case_from_text(text)
    line = manager.get_lines(case.case_id)[0]

    assert line.requested_manufacturer == "Fabrikam Controls"
    assert line.requested_mpn == "ABC-123"
    assert line.quantity == 10
    assert line.review_decision.value == "PENDING"
    assert any(w.code == "LLM_PROPOSED_EXTRACTION" for w in line.extraction_warnings)
    assert case.llm_model == "gpt-6-luna"
    assert case.llm_schema_version == "rfq-grounded-extraction-v1"
    assert case.llm_validator_version == "json-numeric-verbatim-locator-v1"
    assert case.llm_call_count == 2
    assert case.llm_review_call_count == 1


def test_grounded_llm_rows_can_enter_the_review_sample_after_exact_matching(tmp_path):
    text = "Fabrikam Controls ABC-123 10 EA relay"
    catalog = CatalogRepository()
    catalog.add_item(CatalogItem(
        internal_sku="SYN-ABC-123",
        manufacturer="Fabrikam Controls",
        manufacturer_part_number="ABC-123",
        description="Synthetic relay",
    ))
    client = ScriptedClient([response(quote=text, segment_id="draft-line-1"), review_response()])
    manager = CaseManager(
        catalog=catalog,
        storage=DocumentStorage(base_dir=tmp_path),
        llm_extractor=GroundedLLMExtractor(client),
    )

    case = manager.create_case_from_text(text)
    line = manager.get_lines(case.case_id)[0]
    plan = manager.get_review_plan(case.case_id)

    assert line.extraction_status.value == "EXTRACTED"
    assert line.match_status.value == "EXACT_MATCH"
    assert plan.sample_line_ids == [line.line_id]
    assert plan.exception_line_ids == []


def test_repaired_llm_output_is_routed_as_an_exception_not_a_clean_sample(tmp_path):
    text = "Fabrikam Controls ABC-123 10 EA relay"
    catalog = CatalogRepository()
    catalog.add_item(CatalogItem(
        internal_sku="SYN-ABC-123",
        manufacturer="Fabrikam Controls",
        manufacturer_part_number="ABC-123",
        description="Synthetic relay",
    ))
    client = ScriptedClient([
        "not-json", response(quote=text, segment_id="draft-line-1"), review_response()
    ])
    manager = CaseManager(
        catalog=catalog,
        storage=DocumentStorage(base_dir=tmp_path),
        llm_extractor=GroundedLLMExtractor(client),
    )

    case = manager.create_case_from_text(text)
    line = manager.get_lines(case.case_id)[0]
    plan = manager.get_review_plan(case.case_id)

    assert any(w.code == "LLM_REPAIR_USED" for w in line.extraction_warnings)
    assert line.line_id in plan.exception_line_ids
    assert line.line_id not in plan.sample_line_ids


def test_llm_validation_conflict_routes_row_to_exception(tmp_path):
    text = "Fabrikam Controls ABC-123 10 EA relay"
    catalog = CatalogRepository()
    catalog.add_item(CatalogItem(
        internal_sku="SYN-ABC-123",
        manufacturer="Fabrikam Controls",
        manufacturer_part_number="ABC-123",
        description="Synthetic relay",
    ))
    client = ScriptedClient([
        response(quote=text, segment_id="draft-line-1"), review_response("conflict")
    ])
    manager = CaseManager(catalog=catalog, storage=DocumentStorage(base_dir=tmp_path),
                          llm_extractor=GroundedLLMExtractor(client))

    case = manager.create_case_from_text(text)
    line = manager.get_lines(case.case_id)[0]
    plan = manager.get_review_plan(case.case_id)

    assert any(w.code == "LLM_VALIDATION_CONFLICT" for w in line.extraction_warnings)
    assert line.line_id in plan.exception_line_ids
    assert line.line_id not in plan.sample_line_ids


def test_unavailable_llm_validation_routes_row_to_exception(tmp_path):
    text = "Fabrikam Controls ABC-123 10 EA relay"
    client = ScriptedClient([response(quote=text, segment_id="draft-line-1")])
    manager = CaseManager(catalog=CatalogRepository(), storage=DocumentStorage(base_dir=tmp_path),
                          llm_extractor=GroundedLLMExtractor(client))

    case = manager.create_case_from_text(text)
    line = manager.get_lines(case.case_id)[0]

    assert any(w.code == "LLM_VALIDATION_UNAVAILABLE" for w in line.extraction_warnings)
    assert line.line_id in manager.get_review_plan(case.case_id).exception_line_ids


def test_ocr_page_evidence_reaches_llm_fallback_and_remains_exceptional(tmp_path, monkeypatch):
    from rfq_assistant.cases import manager as manager_module
    from rfq_assistant.schemas.models import ExtractionWarning

    quote = "Fabrikam Controls ABC-123 10 EA relay"
    ocr_evidence = EvidenceSegment(
        segment_id="page-1",
        locator=SourceLocator(file_name="scan.pdf", page_number=1),
        text=quote,
        warnings=[ExtractionWarning(
            code="OCR_TEXT_REVIEW", message="Verify OCR against source image.", severity="warning"
        )],
    )

    class FakeParser:
        cache_dependencies = ()

        def cache_profile(self):
            return {"parser": "test.fake", "version": "v1"}

        def parse(self, _content, _filename):
            return ParsedDocumentResult(lines=[], evidence_segments=[ocr_evidence])

    monkeypatch.setattr(manager_module, "get_parser", lambda _source_type: FakeParser())
    client = ScriptedClient([
        response(quote=quote, segment_id="page-1"), review_response()
    ])
    manager = CaseManager(
        catalog=CatalogRepository(),
        storage=DocumentStorage(base_dir=tmp_path),
        llm_extractor=GroundedLLMExtractor(client),
    )

    case = manager.create_case_from_upload("scan.pdf", b"%PDF-synthetic")
    line = manager.get_lines(case.case_id)[0]
    plan = manager.get_review_plan(case.case_id)

    assert line.source_locator.page_number == 1
    assert any(w.code == "OCR_TEXT_REVIEW" for w in line.extraction_warnings)
    assert line.line_id in plan.exception_line_ids


def test_pinned_llm_result_is_reused_after_worker_interruption(tmp_path, monkeypatch):
    from rfq_assistant.cases import manager as manager_module
    from rfq_assistant.persistence import SQLiteStateStore

    quote = "Fabrikam Controls ABC-123 10 EA relay"
    evidence = EvidenceSegment(
        segment_id="row-1",
        locator=SourceLocator(file_name="request.txt", row_number=1),
        text=quote,
    )

    class FakeParser:
        cache_dependencies = ()
        calls = 0

        def cache_profile(self):
            return {"parser": "test.llm-cache", "version": "v1"}

        def parse(self, _content, _filename):
            self.calls += 1
            return ParsedDocumentResult(lines=[], evidence_segments=[evidence])

    parser = FakeParser()
    monkeypatch.setattr(manager_module, "get_parser", lambda *_args, **_kwargs: parser)
    client = ScriptedClient([response(quote=quote), review_response()])
    client.deployment_version = "immutable-test-deployment-v1"
    store = SQLiteStateStore(tmp_path / "state" / "rfq.sqlite3")
    manager = CaseManager(
        catalog=CatalogRepository(),
        storage=DocumentStorage(base_dir=tmp_path / "files"),
        state_store=store,
        llm_extractor=GroundedLLMExtractor(client),
    )
    content = quote.encode("utf-8")
    case = manager.enqueue_case_from_upload("request.txt", content)
    job = store.claim_ingestion_job("worker-1")
    assert job is not None

    original_match = manager_module.ExactMatchEngine.match_line
    fail_once = {"value": True}

    def fail_first_match(matcher, *args, **kwargs):
        if fail_once["value"]:
            fail_once["value"] = False
            raise RuntimeError("synthetic interruption after model validation")
        return original_match(matcher, *args, **kwargs)

    monkeypatch.setattr(manager_module.ExactMatchEngine, "match_line", fail_first_match)
    try:
        manager.create_case_from_upload(
            "request.txt", content, _existing_case_id=case.case_id,
            _job_id=job["job_id"], _worker_id="worker-1",
        )
    except RuntimeError as exc:
        assert "after model validation" in str(exc)
    else:
        raise AssertionError("synthetic interruption was not raised")

    completed = manager.create_case_from_upload(
        "request.txt", content, _existing_case_id=case.case_id,
        _job_id=job["job_id"], _worker_id="worker-1",
    )

    assert len(client.calls) == 2
    assert parser.calls == 1
    assert completed.llm_cache_hit is True
    assert completed.llm_call_count == 2
    assert completed.llm_deployment_version == "immutable-test-deployment-v1"
    assert len(manager.get_lines(case.case_id)) == 1


def test_environment_adapter_is_disabled_by_default(monkeypatch):
    monkeypatch.delenv("RFQ_LLM_ENABLED", raising=False)
    monkeypatch.setenv("RFQ_LLM_API_KEY", "must-not-be-read-or-printed")
    assert GroundedLLMExtractor.from_environment() is None


def test_environment_adapter_accepts_only_luna_and_requires_tls(monkeypatch):
    monkeypatch.setenv("RFQ_LLM_ENABLED", "true")
    monkeypatch.setenv("RFQ_LLM_API_KEY", "synthetic-secret-for-test")
    monkeypatch.setenv("RFQ_LLM_MODEL", "gpt-6-luna")
    monkeypatch.setenv("RFQ_LLM_BASE_URL", "https://model.example.test/v1")
    extractor = GroundedLLMExtractor.from_environment()
    assert extractor is not None
    extractor.close()

    monkeypatch.setenv("RFQ_LLM_MODEL", "some-other-model")
    assert GroundedLLMExtractor.from_environment() is None
    monkeypatch.setenv("RFQ_LLM_MODEL", "not-luna")
    assert GroundedLLMExtractor.from_environment() is None

    monkeypatch.setenv("RFQ_LLM_MODEL", "gpt-6-luna")
    monkeypatch.setenv("RFQ_LLM_BASE_URL", "http://remote.example.test/v1")
    assert GroundedLLMExtractor.from_environment() is None


def test_provider_telemetry_records_latency_and_tokens_without_source_or_secret(caplog):
    import httpx
    caplog.set_level("INFO")

    observed = {}

    def respond(request):
        observed["url"] = str(request.url)
        observed["headers"] = dict(request.headers)
        observed["body"] = request.read().decode("utf-8")
        return httpx.Response(200, json={
            "choices": [{"message": {"content": "{\"lines\":[]}"}}],
            "usage": {"prompt_tokens": 321, "completion_tokens": 11},
        })

    client = OpenAICompatibleClient(
        api_key="synthetic-secret",
        model="gpt-6-luna",
        base_url="https://example.openai.azure.com",
        transport=httpx.MockTransport(respond),
    )
    try:
        result = client.complete(system="synthetic system", user="PRIVATE RFQ SOURCE TEXT")
    finally:
        client.close()

    assert result == '{"lines":[]}'
    assert observed["url"] == "https://example.openai.azure.com/openai/v1/chat/completions"
    assert observed["headers"]["api-key"] == "synthetic-secret"
    assert "temperature" not in json.loads(observed["body"])
    assert "input_tokens=321" in caplog.text
    assert "output_tokens=11" in caplog.text
    assert client.last_usage == {"input_tokens": 321, "output_tokens": 11}
    assert "synthetic-secret" not in caplog.text
    assert "PRIVATE RFQ SOURCE TEXT" not in caplog.text


@pytest.mark.parametrize(
    "first_response,expected_delay",
    [
        (503, 0.2),
        (429, 0.4),
    ],
)
def test_provider_retries_transient_status_once_with_bounded_backoff(first_response, expected_delay):
    import httpx

    attempts = []
    delays = []

    def respond(request):
        attempts.append(request)
        if len(attempts) == 1:
            headers = {"retry-after": "0.4"} if first_response == 429 else None
            return httpx.Response(first_response, headers=headers, json={"error": "temporary"})
        return httpx.Response(200, json={"choices": [{"message": {"content": "ok"}}]})

    client = OpenAICompatibleClient(
        "synthetic-secret", "gpt-6-luna", "https://example.invalid",
        transport=httpx.MockTransport(respond), sleeper=delays.append,
    )
    try:
        assert client.complete(system="system", user="synthetic") == "ok"
        assert client.last_http_attempts == 2
        assert len(attempts) == 2
        assert delays == [expected_delay]
    finally:
        client.close()


def test_provider_retries_transport_timeout_once_without_real_sleep():
    import httpx

    attempts = 0
    delays = []

    def respond(request):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise httpx.ReadTimeout("synthetic timeout", request=request)
        return httpx.Response(200, json={"choices": [{"message": {"content": "ok"}}]})

    client = OpenAICompatibleClient(
        "synthetic-secret", "gpt-6-luna", "https://example.invalid",
        transport=httpx.MockTransport(respond), sleeper=delays.append,
    )
    try:
        assert client.complete(system="system", user="synthetic") == "ok"
        assert attempts == 2
        assert delays == [0.2]
    finally:
        client.close()


@pytest.mark.parametrize(
    "response",
    [
        (400, {}),
        (429, {"retry-after": "30"}),
    ],
)
def test_provider_does_not_retry_permanent_or_over_budget_errors(response):
    import httpx

    status, headers = response
    attempts = []

    def respond(request):
        attempts.append(request)
        return httpx.Response(status, headers=headers, json={"error": "synthetic"})

    client = OpenAICompatibleClient(
        "synthetic-secret", "gpt-6-luna", "https://example.invalid",
        transport=httpx.MockTransport(respond), sleeper=lambda _: None,
    )
    try:
        with pytest.raises(httpx.HTTPStatusError):
            client.complete(system="system", user="synthetic")
        assert len(attempts) == 1
    finally:
        client.close()

def test_source_segment_limit_prevents_accidental_large_model_upload():
    client = ScriptedClient([])
    extractor = GroundedLLMExtractor(client, max_input_characters=5)
    result = extractor.extract([segment()])

    assert result.lines == []
    assert result.warnings[0].code == "LLM_EXTRACTION_BLOCKED"
    assert client.calls == []


def test_no_empty_llm_review_call_is_made_when_extraction_has_no_lines():
    client = ScriptedClient([json.dumps({"lines": []})])
    result = GroundedLLMExtractor(client).extract([segment()])

    assert result.lines == []
    assert len(client.calls) == 1
    assert result.llm_call_count == 1
    assert result.llm_review_call_count == 0


def test_source_segment_builder_fails_closed_before_returning_partial_input():
    from rfq_assistant.schemas.enums import SourceType

    content = ("item,quantity\n" + "\n".join(f"part-{index},1" for index in range(600))).encode()
    try:
        source_segments(content, "large.csv", SourceType.CSV)
    except SourceSegmentLimitError as exc:
        assert "bounded LLM fallback" in str(exc)
    else:
        raise AssertionError("oversized input should be rejected before prompt construction")


def test_xlsx_model_fallback_includes_hidden_rows_with_review_warning():
    import io
    import openpyxl
    from rfq_assistant.schemas.enums import SourceType

    workbook = openpyxl.Workbook()
    visible = workbook.active
    visible.title = "BOM"
    visible.append(["Part Number", "Quantity"])
    visible.append(["VISIBLE-1", 1])
    hidden = workbook.create_sheet("Hidden BOM")
    hidden.sheet_state = "hidden"
    hidden.append(["HiddenPart-2", 4])
    output = io.BytesIO()
    workbook.save(output)

    segments = source_segments(output.getvalue(), "hidden.xlsx", SourceType.XLSX)
    hidden_segment = next(item for item in segments if item.locator.sheet_name == "Hidden BOM")
    assert "HiddenPart-2" in hidden_segment.text
    assert any(w.code == "HIDDEN_SHEET_REVIEW" for w in hidden_segment.warnings)


def test_llm_fallback_recovers_candidate_region_missed_beside_structured_rows(tmp_path):
    text = (
        "Schneider Electric LC1D09BD 2 EA contactor\n"
        "Need quantity: 3; Fabrikam Controls AB 1234 relay"
    )
    candidate_quote = "Need quantity: 3; Fabrikam Controls AB 1234 relay"
    candidate_response = json.dumps({"lines": [{
        "line_number": 1,
        "requested_manufacturer": "Fabrikam Controls",
        "requested_mpn": "AB 1234",
        "description": "relay",
        "quantity": 3,
        "unit": None,
        "segment_id": "candidate-text-2",
        "source_quote": candidate_quote,
    }]})
    catalog = CatalogRepository()
    catalog.add_item(CatalogItem(
        internal_sku="SYN-AB1234",
        manufacturer="Fabrikam Controls",
        manufacturer_part_number="AB 1234",
        description="Synthetic relay",
    ))
    client = ScriptedClient([candidate_response, review_response()])
    manager = CaseManager(
        catalog=catalog,
        storage=DocumentStorage(base_dir=tmp_path),
        llm_extractor=GroundedLLMExtractor(client),
    )

    case = manager.create_case_from_text(text)
    lines = manager.get_lines(case.case_id)

    assert len(lines) == 2
    recovered = next(line for line in lines if line.requested_mpn == "AB 1234")
    assert recovered.line_number == 2
    assert recovered.quantity == 3
    assert recovered.source_locator.row_number == 2
    assert any(w.code == "LLM_PROPOSED_EXTRACTION" for w in recovered.extraction_warnings)
    assert len(client.calls) == 2


def test_unclassified_text_candidate_becomes_visible_manual_review_row(tmp_path):
    text = "Schneider Electric LC1D09BD 2 EA contactor\nNeed special disconnect switches"
    manager = CaseManager(CatalogRepository(), storage=DocumentStorage(base_dir=tmp_path))

    case = manager.create_case_from_text(text)
    lines = manager.get_lines(case.case_id)

    assert len(lines) == 2
    unresolved = next(line for line in lines if "Need special" in line.raw_text)
    assert unresolved.quantity is None
    assert any(w.code == "SOURCE_REGION_UNCLASSIFIED" for w in unresolved.extraction_warnings)
    assert unresolved.line_id in manager.get_review_plan(case.case_id).exception_line_ids


def test_unfamiliar_text_is_sent_to_grounded_fallback_and_can_be_recovered(tmp_path):
    text = "Please quote our bespoke panel component, 4 EA, described as ABC-123 by Fabrikam Controls"
    quote = "our bespoke panel component, 4 EA, described as ABC-123 by Fabrikam Controls"
    proposal = response(
        quantity=4,
        manufacturer="Fabrikam Controls",
        description="bespoke panel component",
        quote=quote,
        segment_id="draft-line-1",
    )
    client = ScriptedClient([proposal, review_response()])
    manager = CaseManager(
        CatalogRepository(),
        storage=DocumentStorage(base_dir=tmp_path),
        llm_extractor=GroundedLLMExtractor(client),
    )

    case = manager.create_case_from_text(text)
    lines = manager.get_lines(case.case_id)

    assert case.status.value != "FAILED"
    assert len(lines) == 1
    assert lines[0].requested_mpn == "ABC-123"
    assert lines[0].requested_manufacturer == "Fabrikam Controls"
    assert lines[0].quantity == 4
    assert len(client.calls) == 2


def test_candidate_only_xlsx_is_retained_for_manual_review(tmp_path):
    import io
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Unmapped"
    sheet.append(["Please quote the custom assembly for our remote panel"])
    output = io.BytesIO()
    workbook.save(output)

    manager = CaseManager(CatalogRepository(), storage=DocumentStorage(base_dir=tmp_path))
    case = manager.create_case_from_upload("unmapped.xlsx", output.getvalue())
    lines = manager.get_lines(case.case_id)

    assert case.status.value != "FAILED"
    assert len(lines) == 1
    assert "custom assembly" in lines[0].raw_text
    assert any(w.code == "SOURCE_REGION_UNCLASSIFIED" for w in lines[0].extraction_warnings)


def test_invalid_model_fallback_preserves_structured_rows_and_candidate_exception(tmp_path):
    text = "Schneider Electric LC1D09BD 2 EA contactor\nNeed special disconnect switches"
    manager = CaseManager(
        CatalogRepository(),
        storage=DocumentStorage(base_dir=tmp_path),
        llm_extractor=GroundedLLMExtractor(ScriptedClient(["not-json", "still-not-json"])),
    )

    case = manager.create_case_from_text(text)
    lines = manager.get_lines(case.case_id)

    assert case.status.value != "FAILED"
    assert any(line.requested_mpn == "LC1D09BD" for line in lines)
    unresolved = next(line for line in lines if "Need special" in line.raw_text)
    assert any(w.code == "SOURCE_REGION_UNCLASSIFIED" for w in unresolved.extraction_warnings)
    assert any(w.code == "LLM_EXTRACTION_BLOCKED" for w in unresolved.extraction_warnings)


def test_xlsx_unknown_header_sheet_is_retained_as_candidate_rows():
    import io
    import openpyxl
    from rfq_assistant.ingestion.parsers.spreadsheet_parser import SpreadsheetParser

    workbook = openpyxl.Workbook()
    known = workbook.active
    known.title = "BOM"
    known.append(["Part Number", "Quantity", "Description"])
    known.append(["LC1D09BD", 2, "contactor"])
    unknown = workbook.create_sheet("Other Items")
    unknown.append(["disconnect assemblies"])
    unknown.append(["auxiliary contacts"])
    output = io.BytesIO()
    workbook.save(output)

    result = SpreadsheetParser().parse(output.getvalue(), "multi_sheet.xlsx")

    assert len(result.lines) == 1
    assert [segment.locator.row_number for segment in result.candidate_segments] == [1, 2]
    assert all(segment.locator.sheet_name == "Other Items" for segment in result.candidate_segments)
