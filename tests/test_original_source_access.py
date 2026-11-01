"""Synthetic source-file access tests for the reviewer workbench."""

import importlib
from io import BytesIO
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

from rfq_assistant.cases.manager import CaseManager
from rfq_assistant.ingestion.storage import DocumentStorage
from rfq_assistant.matching.catalog import CatalogRepository


def test_original_source_is_case_scoped_and_not_cached(tmp_path, monkeypatch):
    api = importlib.import_module("rfq_assistant.api.app")
    manager = CaseManager(catalog=CatalogRepository(), storage=DocumentStorage(tmp_path / "source"))
    monkeypatch.setattr(api, "case_manager", manager)
    monkeypatch.delenv("REQUIRE_AUTH", raising=False)

    original = b"10 pcs Schneider Electric LC1D09BD"
    case = manager.create_case_from_upload(filename="source.txt", content=original)
    client = TestClient(api.app)

    response = client.get(f"/api/rfqs/{case.case_id}/source")
    assert response.status_code == 200
    assert response.content == original
    assert "text/plain" in response.headers["content-type"]
    assert response.headers["cache-control"] == "private, no-store"
    assert response.headers["x-content-type-options"] == "nosniff"

    assert client.get("/api/rfqs/RFQ-NOT-EXIST/source").status_code == 404
    manager.delete_case(case.case_id)
    assert client.get(f"/api/rfqs/{case.case_id}/source").status_code == 404


def test_protected_source_rejects_anonymous_and_allows_viewer(tmp_path, monkeypatch):
    api = importlib.import_module("rfq_assistant.api.app")
    manager = CaseManager(catalog=CatalogRepository(), storage=DocumentStorage(tmp_path / "source"))
    monkeypatch.setattr(api, "case_manager", manager)
    monkeypatch.setenv("REQUIRE_AUTH", "true")
    monkeypatch.setenv("PILOT_API_KEYS", "synthetic-test-viewer-secret-001:viewer:Viewer")
    case = manager.create_case_from_text("10 pcs Schneider Electric LC1D09BD")
    client = TestClient(api.app)

    assert client.get(f"/api/rfqs/{case.case_id}/source").status_code == 401
    response = client.get(
        f"/api/rfqs/{case.case_id}/source",
        headers={"X-API-Key": "synthetic-test-viewer-secret-001"},
    )
    assert response.status_code == 200
    assert b"LC1D09BD" in response.content


def test_spreadsheet_preview_navigates_to_row_and_merged_anchor_cell(tmp_path, monkeypatch):
    api = importlib.import_module("rfq_assistant.api.app")
    manager = CaseManager(catalog=CatalogRepository(), storage=DocumentStorage(tmp_path / "source"))
    monkeypatch.setattr(api, "case_manager", manager)
    monkeypatch.delenv("REQUIRE_AUTH", raising=False)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "RFQ"
    sheet.append(["Manufacturer", "MPN", "Quantity"])
    sheet.append(["Phoenix Contact", "2900364", 2])
    sheet.append([None, "2900365", 4])
    sheet["D3"] = "=1+1"
    sheet.merge_cells("A2:A3")
    output = BytesIO()
    workbook.save(output)
    case = manager.create_case_from_upload(filename="source.xlsx", content=output.getvalue())

    response = TestClient(api.app).get(
        f"/api/rfqs/{case.case_id}/source/preview",
        params={"sheet_name": "RFQ", "row_number": 3, "target_cells": ["A2", "B3"]},
    )
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["sheet_name"] == "RFQ"
    assert preview["target_row"] == 3
    assert preview["target_cells"] == ["A2", "B3"]
    target = next(row for row in preview["rows"] if row["row_number"] == 3)
    assert target["is_target"] is True
    cells = {cell["coordinate"]: cell for cell in target["cells"]}
    assert cells["B3"]["value"] == "2900365"
    assert cells["B3"]["selected"] is True
    assert cells["D3"]["is_formula"] is True
    assert cells["D3"]["value"] == "=1+1"
    anchor = next(row for row in preview["rows"] if row["row_number"] == 2)
    assert next(cell for cell in anchor["cells"] if cell["coordinate"] == "A2")["selected"] is True


def test_spreadsheet_preview_bounds_rows_and_rejects_unknown_sheets(tmp_path, monkeypatch):
    api = importlib.import_module("rfq_assistant.api.app")
    manager = CaseManager(catalog=CatalogRepository(), storage=DocumentStorage(tmp_path / "source"))
    monkeypatch.setattr(api, "case_manager", manager)
    monkeypatch.delenv("REQUIRE_AUTH", raising=False)
    workbook = openpyxl.Workbook()
    workbook.active.title = "RFQ"
    workbook.active.append(["MPN", "Quantity"])
    workbook.active.append(["PN-0001", 1])
    output = BytesIO()
    workbook.save(output)
    case = manager.create_case_from_upload(filename="source.xlsx", content=output.getvalue())
    client = TestClient(api.app)

    base_url = f"/api/rfqs/{case.case_id}/source/preview"
    assert client.get(base_url, params={"sheet_name": "Missing", "row_number": 1}).status_code == 404
    assert client.get(base_url, params={"sheet_name": "RFQ", "row_number": 5001}).status_code == 422
    assert client.get(
        base_url,
        params={"sheet_name": "RFQ", "row_number": 1, "target_cells": "A1;DROP"},
    ).status_code == 422


def test_csv_source_preview_highlights_the_recorded_row(tmp_path, monkeypatch):
    api = importlib.import_module("rfq_assistant.api.app")
    manager = CaseManager(catalog=CatalogRepository(), storage=DocumentStorage(tmp_path / "source"))
    monkeypatch.setattr(api, "case_manager", manager)
    monkeypatch.delenv("REQUIRE_AUTH", raising=False)
    case = manager.create_case_from_upload(
        filename="source.csv",
        content=b"Manufacturer,MPN,Quantity\nPhoenix Contact,2900364,2\nPhoenix Contact,2900365,4\n",
    )

    response = TestClient(api.app).get(
        f"/api/rfqs/{case.case_id}/source/preview",
        params={"sheet_name": "Default", "row_number": 2},
    )
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["sheet_name"] == "CSV"
    assert preview["target_row"] == 2
    target = next(row for row in preview["rows"] if row["row_number"] == 2)
    assert target["is_target"] is True
    assert [cell["value"] for cell in target["cells"]] == ["Phoenix Contact", "2900364", "2"]
    assert all(cell["selected"] for cell in target["cells"])


def test_pdf_page_preview_renders_only_the_case_scoped_page(tmp_path, monkeypatch):
    api = importlib.import_module("rfq_assistant.api.app")
    manager = CaseManager(catalog=CatalogRepository(), storage=DocumentStorage(tmp_path / "source"))
    monkeypatch.setattr(api, "case_manager", manager)
    monkeypatch.delenv("REQUIRE_AUTH", raising=False)
    fixture = Path(__file__).resolve().parent.parent / "fixtures" / "rfq_samples" / "formal_rfq_native.pdf"
    case = manager.create_case_from_upload(filename="source.pdf", content=fixture.read_bytes())

    client = TestClient(api.app)
    response = client.get(f"/api/rfqs/{case.case_id}/source/pages/1")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "image/png"
    assert response.content.startswith(b"\x89PNG\r\n\x1a\n")
    assert response.headers["cache-control"] == "private, no-store"
    assert client.get(f"/api/rfqs/{case.case_id}/source/pages/999").status_code == 404
    assert client.get("/api/rfqs/RFQ-NOT-EXIST/source/pages/1").status_code == 404
