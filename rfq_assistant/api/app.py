"""FastAPI application for RFQ Assistant backend API."""

import hashlib
import json
import logging
import mimetypes
import os
import time
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import List, Literal, Optional

from fastapi import (
    Body,
    Depends,
    FastAPI,
    File,
    Form,
    HTTPException,
    Query,
    Request,
    UploadFile,
    status,
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response

from rfq_assistant.agents.tools import ApprovalGateError, QuotationCommercialService
from rfq_assistant.api.cors import configured_cors_origins
from rfq_assistant.api.body_limit import MultipartBodyLimitMiddleware
from rfq_assistant.api.intake_limit import IntakeRateLimiter
from rfq_assistant.auth import (
    PilotUser,
    env_flag,
    get_current_user,
    initialize_authentication,
    require_role,
)
from rfq_assistant.cases.manager import CaseManager
from rfq_assistant.exports import ExportGenerator, ExportValidationError
from rfq_assistant.ingestion.ai_extraction import GroundedLLMExtractor
from rfq_assistant.ingestion.excel import ExcelBOMInspector
from rfq_assistant.ingestion.excel_values import render_excel_cell
from rfq_assistant.ingestion.malware import ClamAVScanner
from rfq_assistant.ingestion.parsers.base import DocumentParserError
from rfq_assistant.ingestion.storage import DocumentStorage
from rfq_assistant.ingestion.validator import MAX_UPLOAD_BYTES, UploadValidationError
from rfq_assistant.ingestion.worker import RFQIngestionWorker
from rfq_assistant.knowledge.catalog_updates import (
    CatalogUpdateError,
    CatalogUpdateService,
)
from rfq_assistant.matching.catalog import CatalogImportError, CatalogRepository
from rfq_assistant.matching.fuzzy import RapidFuzzCandidateEngine
from rfq_assistant.persistence import SQLiteStateStore
from rfq_assistant.quotation.proposal import (
    GroundedQuoteProposalService,
    QuoteProposalError,
)
from rfq_assistant.quotation.validate import (
    ProposedQuote,
    QuoteLine,
    SourceRef,
    calculate_quote_total,
    validate_proposed_quote,
)
from rfq_assistant.schemas.enums import AuditAction
from rfq_assistant.schemas.models import (
    ApproveRemainingRequest,
    AuditEvent,
    BatchReviewRequest,
    CatalogItem,
    CatalogUpdateProposal,
    CatalogUpdateReviewRequest,
    ExportDocumentSummary,
    ExportManifest,
    ExportRequest,
    ExtractionWarning,
    HealthResponse,
    RematchLineRequest,
    ReviewPlan,
    RFQCase,
    RFQCaseSummary,
    RFQLineItem,
    RFQLineItemUpdateRequest,
    SourcePreviewCell,
    SourcePreviewRow,
    SpreadsheetSourcePreview,
)

ROOT = Path(__file__).resolve().parent.parent.parent
log = logging.getLogger(__name__)
DEFAULT_CATALOG_PATH = ROOT / "fixtures" / "catalogs" / "industrial_components_sample_catalog.csv"
MAX_CATALOG_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_RFQ_SOURCES_PER_CASE = 6
MAX_RFQ_SOURCE_BYTES_PER_CASE = 25 * 1024 * 1024
MAX_MULTIPART_BODY_BYTES = MAX_RFQ_SOURCE_BYTES_PER_CASE + 1024 * 1024


class PNGResponse(Response):
    media_type = "image/png"

# Global repositories and services
_storage_root = Path(os.environ.get("RFQ_ASSISTANT_STORAGE_DIR", "storage"))
_database_path = Path(os.environ.get("RFQ_ASSISTANT_DB_PATH", str(_storage_root / "rfq_assistant.sqlite3")))
state_store = SQLiteStateStore(_database_path)
catalog_repo = CatalogRepository(state_store=state_store)
case_manager = CaseManager(
    catalog=catalog_repo,
    storage=DocumentStorage(base_dir=_storage_root),
    state_store=state_store,
    llm_extractor=GroundedLLMExtractor.from_environment(),
)
export_generator = ExportGenerator(catalog=catalog_repo, base_export_dir=_storage_root / "exports")
catalog_update_service = CatalogUpdateService(
    catalog=catalog_repo,
    state_store=state_store,
    client=case_manager.llm_extractor.client if case_manager.llm_extractor else None,
)
catalog_update_source_storage = DocumentStorage(base_dir=_storage_root / "distributor_sources")
clamav_scanner = ClamAVScanner()
rapidfuzz_engine = RapidFuzzCandidateEngine()
commercial_service = QuotationCommercialService(state_store=state_store)
ASYNC_INGEST_ENABLED = env_flag("RFQ_ASYNC_INGEST_ENABLED", default=True)
try:
    INGEST_WORKER_CONCURRENCY = max(1, min(4, int(os.getenv("RFQ_INGEST_WORKER_CONCURRENCY", "2"))))
except ValueError:
    INGEST_WORKER_CONCURRENCY = 2
    log.warning("Invalid RFQ_INGEST_WORKER_CONCURRENCY; using the safe default of 2.")
ingestion_worker = (
    RFQIngestionWorker(case_manager, state_store, max_concurrent_jobs=INGEST_WORKER_CONCURRENCY)
    if ASYNC_INGEST_ENABLED else None
)

# Pre-load default synthetic catalog if available
if DEFAULT_CATALOG_PATH.exists() and catalog_repo.count() == 0:
    try:
        catalog_repo.import_csv(DEFAULT_CATALOG_PATH.read_bytes(), version="v1-synth")
        rapidfuzz_engine.index_catalog(catalog_repo.list_all())
    except Exception as e:
        print(f"Warning: Failed to pre-load default catalog: {e}")

def _rfq_source_snapshot_hash(case_id: str) -> str:
    """Fingerprint the *current* RFQ and item master, not a saved quotation copy.

    Including the full prototype catalog is deliberately conservative: any
    import/update requires re-drafting and re-approval. Customer deployments
    should replace this with versioned, workspace-scoped catalog snapshots.
    """
    case = case_manager.get_case(case_id)
    if case is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Source RFQ no longer exists. Quote must not be approved or sent.",
        )
    payload = {
        "case_id": case_id,
        "customer_label": case.customer_label,
        "file_digest": case.source_file.sha256_digest if case.source_file else None,
        "lines": [line.model_dump(mode="json") for line in case_manager.get_lines(case_id)],
        "catalog": [
            item.model_dump(mode="json")
            for item in sorted(catalog_repo.list_all(), key=lambda row: row.internal_sku)
        ],
    }
    canonical = json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def _require_fresh_quote_source(quote_id: str) -> None:
    """Reject a saved draft if its underlying RFQ or catalog changed."""
    draft = commercial_service.get_draft(quote_id)
    if draft is None:
        return  # The commercial service reports the missing quote.
    if not draft.source_state_hash or draft.source_state_hash != _rfq_source_snapshot_hash(draft.case_id):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="RFQ lines, reviewer decisions or catalog changed since this draft was created. Regenerate and review the quote.",
        )


@asynccontextmanager
async def lifespan(_app: FastAPI) -> AsyncIterator[None]:
    initialize_authentication()
    if ingestion_worker:
        ingestion_worker.start()
    try:
        yield
    finally:
        if ingestion_worker:
            ingestion_worker.stop()
        if case_manager.llm_extractor:
            case_manager.llm_extractor.close()


app = FastAPI(
    title="RFQ Assistant API",
    version="0.2.0",
    description="Human-reviewed RFQ intake and quotation preparation assistant with OSS stack",
    lifespan=lifespan,
)

# Starlette parses and may spool multipart files before endpoint-level file
# checks run. Cap the entire wire body first, with 1 MiB reserved for fields and
# multipart framing beyond the 25 MiB source-content budget.
app.add_middleware(
    MultipartBodyLimitMiddleware,
    max_bytes=MAX_MULTIPART_BODY_BYTES,
)

# This caps expensive admission in a single local API process. A real multi-host
# distributor deployment also requires shared rate limits at the trusted edge.
try:
    _intake_rate = int(os.getenv("RFQ_MAX_INTAKES_PER_MINUTE", "240"))
    app.state.intake_rate_limiter = IntakeRateLimiter(_intake_rate)
except ValueError as exc:
    raise RuntimeError("Invalid RFQ_MAX_INTAKES_PER_MINUTE; configure 1..1000.") from exc


@app.middleware("http")
async def cap_expensive_intake(request: Request, call_next):
    # Use the ASGI router's raw path. Starlette versions affected by
    # CVE-2026-48710 could let a malformed Host header rewrite request.url.path
    # while the router still dispatched the original path.
    request_path = request.scope.get("path", "")
    if request.method == "POST" and request_path in {
        "/api/rfqs", "/api/catalogs/import", "/api/catalogs/updates", "/v1/rfqs/extract",
    }:
        allowed, retry_after = request.app.state.intake_rate_limiter.allow()
        if not allowed:
            return JSONResponse(
                status_code=429,
                content={"detail": "Too many RFQ intake requests. Retry shortly."},
                headers={"Retry-After": str(retry_after), "Cache-Control": "no-store"},
            )
    return await call_next(request)

# Enable CORS for explicitly configured browser origins.
_cors_origins = configured_cors_origins()
app.add_middleware(
    CORSMiddleware,
    allow_origins=_cors_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def _scan_blocks_intake(scan) -> bool:
    """Block detected malware and unscanned content in protected deployments."""
    if getattr(scan, "quarantined", False):
        return True
    if not getattr(scan, "scanned", True):
        return clamav_scanner.protected_mode
    return not scan.is_clean


def _unscanned_warning(scan) -> List[ExtractionWarning]:
    if getattr(scan, "scanned", True):
        return []
    return [ExtractionWarning(
        code="ANTIVIRUS_SCAN_UNAVAILABLE",
        message="No antivirus engine scanned this source. Review the original file carefully before using extracted data.",
        severity="warning",
    )]


async def _read_bounded_upload(file: UploadFile) -> bytes:
    """Read at most the configured file limit plus one sentinel byte."""
    if file.size is not None and file.size > MAX_UPLOAD_BYTES:
        await file.close()
        raise HTTPException(
            status_code=status.HTTP_413_CONTENT_TOO_LARGE,
            detail=f"File exceeds the {MAX_UPLOAD_BYTES}-byte upload limit.",
        )
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        await file.close()
        raise HTTPException(
            status_code=status.HTTP_413_CONTENT_TOO_LARGE,
            detail=f"File exceeds the {MAX_UPLOAD_BYTES}-byte upload limit.",
        )
    return content


def _case_source_metadata(case: RFQCase, source_file_id: Optional[str]):
    if source_file_id is None:
        return case.source_file
    for document in case.source_documents:
        if document.source_file.file_id == source_file_id:
            return document.source_file
    return None


@app.get("/api/health", response_model=HealthResponse)
def health_check():
    """Public liveness probe; expose no catalog or customer activity counts."""
    return {"status": "healthy"}


@app.get("/api/session", response_model=dict[str, str])
def get_authenticated_session(current_user: PilotUser = Depends(get_current_user)):
    """Return the identity bound to this API credential, never the credential itself."""
    return {
        "user_id": current_user.user_id,
        "name": current_user.name,
        "role": current_user.role,
        "workspace_id": current_user.workspace_id,
    }


# ---------------------------------------------------------------------------
# Catalog Endpoints
# ---------------------------------------------------------------------------

@app.post("/api/catalogs/import")
async def import_catalog(
    file: UploadFile = File(...),
    current_user: PilotUser = Depends(require_role(["admin"])),
):
    """Imports this deployment's single-workspace catalog CSV (admin only)."""
    content = await file.read(MAX_CATALOG_UPLOAD_BYTES + 1)
    if len(content) > MAX_CATALOG_UPLOAD_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_CONTENT_TOO_LARGE,
            detail=f"Catalog CSV exceeds the {MAX_CATALOG_UPLOAD_BYTES}-byte upload limit.",
        )
    try:
        old_version = catalog_repo.active_version
        summary = catalog_repo.replace_csv(content)
        new_version = summary.get("catalog_version")
        if not isinstance(new_version, str):
            raise HTTPException(status_code=500, detail="Catalog import did not produce a valid catalog version.")
        if new_version != old_version:
            case_manager.handle_catalog_replacement(current_user.name, new_version)
        return {"status": "success", "summary": summary}
    except CatalogImportError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@app.get("/api/catalogs", response_model=List[CatalogItem])
def list_catalog(current_user: PilotUser = Depends(get_current_user)):
    """Lists all catalog items in item master."""
    return catalog_repo.list_all()


@app.get("/api/catalogs/updates", response_model=List[CatalogUpdateProposal])
def list_catalog_updates(current_user: PilotUser = Depends(get_current_user)):
    """List staged and reviewed distributor-data proposals."""
    return catalog_update_service.list_proposals()


@app.get("/api/catalogs/updates/{proposal_id}/source")
def get_catalog_update_source(
    proposal_id: str,
    current_user: PilotUser = Depends(get_current_user),
):
    """Return the retained original source for an authenticated reviewer."""
    proposal = catalog_update_service.get_proposal(proposal_id)
    if proposal is None:
        raise HTTPException(status_code=404, detail="Catalog update proposal not found.")
    suffix = Path(proposal.source_name).suffix.lower() or ".txt"
    source = catalog_update_source_storage.get_file("catalog-updates", f"{proposal.source_sha256}{suffix}")
    if source is None:
        raise HTTPException(status_code=404, detail="The original distributor source is unavailable.")
    source_type = mimetypes.guess_type(f"source{suffix}")[0] or "application/octet-stream"
    return Response(
        content=source,
        media_type=source_type,
        headers={
            "Cache-Control": "private, no-store",
            "Content-Disposition": f'inline; filename="{proposal.proposal_id}{suffix}"',
            "X-Content-Type-Options": "nosniff",
        },
    )


@app.post("/api/catalogs/updates", response_model=List[CatalogUpdateProposal])
async def stage_catalog_update(
    file: Optional[UploadFile] = File(None),
    pasted_text: Optional[str] = Form(None),
    revision_of: Optional[str] = Form(None),
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """Preserve a distributor source and stage grounded proposals without applying them."""
    if file is None and not (pasted_text and pasted_text.strip()):
        raise HTTPException(status_code=400, detail="Paste a distributor note or upload a source file.")
    filename = Path(file.filename or "distributor-update.txt").name if file else "pasted-distributor-update.txt"
    content = await file.read(10 * 1024 * 1024 + 1) if file else (pasted_text or "").encode("utf-8")
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Distributor update source exceeds the 10 MB limit.")
    source_warnings: List[str] = []
    if file:
        scan = clamav_scanner.scan_bytes(content, filename=filename)
        if _scan_blocks_intake(scan):
            raise HTTPException(
                status_code=400,
                detail=("Security scan unavailable. Processing is blocked in protected mode."
                        if not scan.scanned and clamav_scanner.protected_mode
                        else f"Security alert: Malicious attachment quarantined ({scan.threat_name}). Processing blocked."),
            )
        if not scan.scanned:
            source_warnings.append("No antivirus engine scanned this source. Review the original before accepting any changes.")
    digest = hashlib.sha256(content).hexdigest()
    suffix = Path(filename).suffix.lower() or ".txt"
    safe_source_name = f"{digest}{suffix}"
    try:
        catalog_update_source_storage.save_file("catalog-updates", safe_source_name, content)
        proposals = catalog_update_service.stage(
            content=content,
            filename=filename,
            actor=current_user.name,
            pasted_text=pasted_text if file is None else None,
            revision_of=revision_of,
            source_warnings=source_warnings,
        )
        return proposals
    except (CatalogUpdateError, ValueError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.post("/api/catalogs/updates/{proposal_id}/accept", response_model=CatalogUpdateProposal)
def accept_catalog_update(
    proposal_id: str,
    request: CatalogUpdateReviewRequest = Body(default=CatalogUpdateReviewRequest()),
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """Apply one complete, grounded proposal and invalidate stale RFQ matches."""
    try:
        proposal = catalog_update_service.review(proposal_id, actor=current_user.name, accept=True, note=request.note)
        if proposal.current_item and proposal.current_item.catalog_version:
            case_manager.handle_catalog_replacement(current_user.name, proposal.current_item.catalog_version)
        return proposal
    except CatalogUpdateError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


@app.post("/api/catalogs/updates/{proposal_id}/reject", response_model=CatalogUpdateProposal)
def reject_catalog_update(
    proposal_id: str,
    request: CatalogUpdateReviewRequest = Body(default=CatalogUpdateReviewRequest()),
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """Record rejection while leaving active distributor records untouched."""
    try:
        return catalog_update_service.review(proposal_id, actor=current_user.name, accept=False, note=request.note)
    except CatalogUpdateError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


# ---------------------------------------------------------------------------
# RFQ Case Endpoints
# ---------------------------------------------------------------------------

@app.post("/api/rfqs", response_model=RFQCase)
async def create_rfq(
    file: Optional[UploadFile] = File(None),
    files: Optional[List[UploadFile]] = File(None),
    pasted_text: Optional[str] = Form(None),
    title: Optional[str] = Form(None),
    customer_label: Optional[str] = Form(None),
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """Creates draft case from an uploaded file or pasted email text with malware defense."""
    if file and files:
        raise HTTPException(status_code=400, detail="Use either the single-file field or the multiple-files field, not both.")
    attached_files = files or ([file] if file else [])
    if not attached_files and not pasted_text:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Must provide either an uploaded file or pasted text.",
        )

    # Audit identity always comes from the auth dependency. The local demo
    # principal is explicit; client form fields must not create fake actors.
    effective_actor = current_user.name

    try:
        if len(attached_files) + int(bool(pasted_text)) > MAX_RFQ_SOURCES_PER_CASE:
            raise HTTPException(
                status_code=413,
                detail=f"An RFQ can contain at most {MAX_RFQ_SOURCES_PER_CASE} source documents.",
            )
        if pasted_text and len(pasted_text.encode("utf-8")) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail="Pasted RFQ text exceeds the 10 MB source limit.")

        prepared_uploads: List[tuple[str, bytes]] = []
        all_intake_warnings: List[ExtractionWarning] = []
        for attached_file in attached_files:
            read_started = time.perf_counter()
            content = await _read_bounded_upload(attached_file)
            upload_read_ms = (time.perf_counter() - read_started) * 1000
            filename = attached_file.filename or "upload.xlsx"
            if sum(len(existing_content) for _, existing_content in prepared_uploads) + len(content) > MAX_RFQ_SOURCE_BYTES_PER_CASE:
                raise HTTPException(
                    status_code=413,
                    detail=f"Combined RFQ source size exceeds the {MAX_RFQ_SOURCE_BYTES_PER_CASE}-byte limit.",
                )
            # 1. Attachment security scan
            scan_started = time.perf_counter()
            scan = clamav_scanner.scan_bytes(content, filename=filename)
            log.info(
                "[rfq_timing] stage=upload_security bytes=%d upload_read_ms=%.1f scanner_ms=%.1f clean=%s",
                len(content), upload_read_ms, (time.perf_counter() - scan_started) * 1000, scan.is_clean,
            )
            if _scan_blocks_intake(scan):
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=(
                        "Security scan unavailable. Processing is blocked in protected mode."
                        if not scan.scanned and clamav_scanner.protected_mode
                        else f"Security alert: Malicious attachment quarantined ({scan.threat_name}). Processing blocked."
                    ),
                )

            prepared_uploads.append((filename, content))
            all_intake_warnings.extend(_unscanned_warning(scan))

        raw_text = pasted_text or ""
        if raw_text:
            scan_started = time.perf_counter()
            scan = clamav_scanner.scan_bytes(raw_text.encode("utf-8"), filename="pasted_text.txt")
            log.info(
                "[rfq_timing] stage=pasted_text_security scanner_ms=%.1f clean=%s",
                (time.perf_counter() - scan_started) * 1000, scan.is_clean,
            )
            if _scan_blocks_intake(scan):
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=(
                        "Security scan unavailable. Processing is blocked in protected mode."
                        if not scan.scanned and clamav_scanner.protected_mode
                        else f"Security alert: Malicious content quarantined ({scan.threat_name})."
                    ),
                )

            prepared_uploads.insert(0, ("customer_email.txt", raw_text.encode("utf-8")))
            all_intake_warnings.extend(_unscanned_warning(scan))

        if len(prepared_uploads) == 1 and not raw_text:
            filename, content = prepared_uploads[0]
            if ASYNC_INGEST_ENABLED:
                return case_manager.enqueue_case_from_upload(
                    filename=filename, content=content, title=title,
                    customer_label=customer_label, actor=effective_actor,
                    intake_warnings=all_intake_warnings,
                )
            return case_manager.create_case_from_upload(
                filename=filename, content=content, title=title,
                customer_label=customer_label, actor=effective_actor,
                intake_warnings=all_intake_warnings,
            )
        if len(prepared_uploads) == 1 and raw_text:
            if ASYNC_INGEST_ENABLED:
                return case_manager.enqueue_case_from_text(
                    raw_text=raw_text, title=title, customer_label=customer_label,
                    actor=effective_actor, intake_warnings=all_intake_warnings,
                )
            return case_manager.create_case_from_text(
                raw_text=raw_text, title=title, customer_label=customer_label,
                actor=effective_actor, intake_warnings=all_intake_warnings,
            )
        return case_manager.create_case_from_uploads(
            prepared_uploads,
            title=title,
            customer_label=customer_label,
            actor=effective_actor,
            intake_warnings=all_intake_warnings,
            defer_processing=ASYNC_INGEST_ENABLED,
        )
    except (UploadValidationError, DocumentParserError) as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except HTTPException:
        raise
    except Exception as exc:
        log.exception("Unexpected RFQ intake failure")
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Unable to process this RFQ. Verify the file and try again.",
        ) from exc


@app.get("/api/rfqs", response_model=List[RFQCaseSummary])
def list_rfqs(current_user: PilotUser = Depends(get_current_user)):
    """Returns all RFQ cases sorted by date."""
    return case_manager.list_cases()


@app.get("/api/rfqs/{case_id}", response_model=RFQCase)
def get_rfq(case_id: str, current_user: PilotUser = Depends(get_current_user)):
    case = case_manager.get_case(case_id)
    if not case:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")
    return case


@app.get("/api/rfqs/{case_id}/lines", response_model=List[RFQLineItem])
def get_rfq_lines(case_id: str, current_user: PilotUser = Depends(get_current_user)):
    case = case_manager.get_case(case_id)
    if not case:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")
    return case_manager.get_lines(case_id)


@app.get("/api/rfqs/{case_id}/source")
def get_original_source(
    case_id: str,
    source_file_id: Optional[str] = Query(default=None),
    current_user: PilotUser = Depends(get_current_user),
):
    """Open the original uploaded RFQ for human source verification.

    Always resolve the stored file from an existing authorized case, never from
    an arbitrary path supplied in the URL. Native PDF/text can open inline;
    XLSX is downloaded for inspection in the operator's spreadsheet tool.
    """
    from fastapi.responses import FileResponse

    case = case_manager.get_case(case_id)
    source_metadata = _case_source_metadata(case, source_file_id) if case else None
    if not case or not source_metadata:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ source not found.")

    source = case_manager.storage.get_path(
        case.case_id, source_metadata.storage_key or source_metadata.file_name
    )
    if source is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Original source file is missing.")

    ext = Path(source_metadata.file_name).suffix.lower()
    mime = {
        ".pdf": "application/pdf",
        ".csv": "text/csv",
        ".txt": "text/plain",
        ".xls": "application/vnd.ms-excel",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }.get(ext, "application/octet-stream")

    return FileResponse(
        path=source,
        filename=source_metadata.file_name,
        media_type=mime,
        content_disposition_type="inline" if ext in {".pdf", ".txt", ".csv"} else "attachment",
        headers={"Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff"},
    )


@app.post("/api/rfqs/{case_id}/sources/{source_file_id}/retry", response_model=RFQCase)
def retry_rfq_source_document(
    case_id: str,
    source_file_id: str,
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """Retry one failed attachment while preserving other source results."""
    try:
        return case_manager.retry_source_document(case_id, source_file_id, current_user.name)
    except ValueError as exc:
        status_code = 404 if "not found" in str(exc).lower() else 409
        raise HTTPException(status_code=status_code, detail=str(exc)) from exc


@app.get(
    "/api/rfqs/{case_id}/source/pages/{page_number}",
    response_class=PNGResponse,
    responses={200: {"content": {"image/png": {"schema": {"type": "string", "format": "binary"}}}}},
)
def get_pdf_source_page(
    case_id: str,
    page_number: int,
    source_file_id: Optional[str] = Query(default=None),
    current_user: PilotUser = Depends(get_current_user),
):
    """Render one bounded PDF page as a safe image for source verification."""
    import math

    import pymupdf
    if page_number < 1:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Page number must be positive.")
    case = case_manager.get_case(case_id)
    source_metadata = _case_source_metadata(case, source_file_id) if case else None
    if not case or not source_metadata:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ source not found.")
    source = case_manager.storage.get_path(
        case.case_id, source_metadata.storage_key or source_metadata.file_name
    )
    if source is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Original source file is missing.")
    if Path(source_metadata.file_name).suffix.lower() != ".pdf":
        raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Source page previews are available for PDF documents.")

    document = None
    try:
        document = pymupdf.open(source)
        if page_number > document.page_count:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="PDF page not found.")
        page = document.load_page(page_number - 1)
        width, height = page.rect.width, page.rect.height
        if not math.isfinite(width) or not math.isfinite(height) or width <= 0 or height <= 0:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="PDF page dimensions are invalid.")
        scale = min(1.0, 1600 / max(width, height))
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False, annots=False)
        return Response(
            content=pixmap.tobytes("png"),
            media_type="image/png",
            headers={"Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff"},
        )
    except HTTPException:
        raise
    except Exception as exc:
        log.warning("[rfq_pdf_page_preview] failed case=%s page=%d error=%s", case_id, page_number, type(exc).__name__)
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="PDF page could not be previewed safely.") from exc
    finally:
        if document is not None:
            document.close()


@app.get(
    "/api/rfqs/{case_id}/source/preview",
    response_model=SpreadsheetSourcePreview,
)
def get_spreadsheet_source_preview(
    case_id: str,
    sheet_name: str,
    row_number: int,
    target_cells: Optional[List[str]] = Query(default=None),
    radius: int = Query(default=4, ge=0, le=10),
    source_file_id: Optional[str] = Query(default=None),
    current_user: PilotUser = Depends(get_current_user),
):
    """Return a small read-only worksheet window centered on source evidence."""
    import re

    import openpyxl
    from openpyxl.utils.cell import (
        column_index_from_string,
        coordinate_from_string,
        get_column_letter,
    )

    if not 1 <= row_number <= 5000:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Source row is outside workbook limits.")
    case = case_manager.get_case(case_id)
    source_metadata = _case_source_metadata(case, source_file_id) if case else None
    if not case or not source_metadata:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ source not found.")

    source = case_manager.storage.get_path(
        case.case_id, source_metadata.storage_key or source_metadata.file_name
    )
    if source is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Original source file is missing.")
    source_ext = Path(source_metadata.file_name).suffix.lower()
    if source_ext not in {".xlsx", ".csv"}:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Inline source preview is available for CSV and XLSX files. Open the original file to inspect this source.",
        )

    selected_cells: list[str] = []
    for coordinate in target_cells or []:
        if len(selected_cells) >= 20 or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", coordinate, re.IGNORECASE):
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Invalid source cell reference.")
        column, cell_row = coordinate_from_string(coordinate.upper())
        if len(column) > 3 or cell_row > 5000:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Source cell is outside workbook limits.")
        selected_cells.append(coordinate.upper())

    if source_ext == ".csv":
        import csv

        try:
            with source.open("r", encoding="utf-8-sig", newline="") as stream:
                reader = csv.reader(stream, strict=True)
                all_rows = []
                for index, values in enumerate(reader, start=1):
                    if index > 5000 or len(values) > 100:
                        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="CSV exceeds source preview limits.")
                    all_rows.append(values)
            if not 1 <= row_number <= len(all_rows):
                raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Source row is not present in this CSV.")
            for coordinate in selected_cells:
                column_letters, cell_row = coordinate_from_string(coordinate)
                if cell_row > len(all_rows) or column_index_from_string(column_letters) > len(all_rows[cell_row - 1]):
                    raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Source cell is not present in this CSV.")
            rows = []
            for target_row in range(max(1, row_number - radius), min(len(all_rows), row_number + radius) + 1):
                values = all_rows[target_row - 1]
                cells = [
                    SourcePreviewCell(
                        coordinate=f"{get_column_letter(column)}{target_row}",
                        value=value,
                        selected=target_row == row_number and (not selected_cells or f"{get_column_letter(column)}{target_row}".upper() in selected_cells),
                    )
                    for column, value in enumerate(values, start=1)
                ]
                rows.append(SourcePreviewRow(row_number=target_row, is_target=target_row == row_number, cells=cells))
            return SpreadsheetSourcePreview(
                file_name=source_metadata.file_name,
                sheet_name="CSV",
                target_row=row_number,
                target_cells=selected_cells,
                rows=rows,
            )
        except HTTPException:
            raise
        except Exception as exc:
            log.warning("[rfq_csv_source_preview] failed case=%s error=%s", case_id, type(exc).__name__)
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Source CSV could not be previewed safely.") from exc

    workbook = cached_workbook = None
    try:
        source_bytes = source.read_bytes()
        workbook = openpyxl.load_workbook(BytesIO(source_bytes), read_only=True, data_only=False)
        cached_workbook = openpyxl.load_workbook(BytesIO(source_bytes), read_only=True, data_only=True)
        if sheet_name not in workbook.sheetnames:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Worksheet not found in source workbook.")
        worksheet = workbook[sheet_name]
        cached_worksheet = cached_workbook[sheet_name]
        if (worksheet.max_row and worksheet.max_row > 5000) or (worksheet.max_column and worksheet.max_column > 100):
            raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Worksheet exceeds source preview limits.")
        max_row = worksheet.max_row or 0
        max_column = worksheet.max_column or 0
        if not 1 <= row_number <= max_row:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Source row is not present in this worksheet.")
        columns = min(max(worksheet.max_column or 1, 1), 100)
        row_number_set = set(range(max(1, row_number - radius), min(max_row, row_number + radius) + 1))
        for coordinate in selected_cells:
            column_letters, cell_row = coordinate_from_string(coordinate)
            if cell_row > max_row or column_index_from_string(column_letters) > max_column:
                raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Source cell is not present in this worksheet.")
            row_number_set.add(cell_row)
        row_numbers = sorted(row_number_set)
        rows = []
        for target_row in row_numbers:
            formula_cells = next(worksheet.iter_rows(min_row=target_row, max_row=target_row, min_col=1, max_col=columns))
            cached_cells = next(cached_worksheet.iter_rows(min_row=target_row, max_row=target_row, min_col=1, max_col=columns))
            cells = []
            for index, cell in enumerate(formula_cells, start=1):
                coordinate = f"{get_column_letter(index)}{target_row}"
                if cell.value is None and coordinate.upper() not in selected_cells:
                    continue
                cells.append(SourcePreviewCell(
                    coordinate=coordinate,
                    value=render_excel_cell(cell, cached_cells[index - 1]),
                    is_formula=cell.data_type == "f",
                    selected=coordinate.upper() in selected_cells,
                ))
            rows.append(SourcePreviewRow(row_number=target_row, is_target=target_row == row_number, cells=cells))
        return SpreadsheetSourcePreview(
            file_name=source_metadata.file_name,
            sheet_name=sheet_name,
            target_row=row_number,
            target_cells=selected_cells,
            rows=rows,
        )
    except HTTPException:
        raise
    except Exception as exc:
        log.warning("[rfq_source_preview] failed case=%s error=%s", case_id, type(exc).__name__, exc_info=True)
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Source worksheet could not be previewed safely.") from exc
    finally:
        if workbook is not None:
            workbook.close()
        if cached_workbook is not None:
            cached_workbook.close()


@app.patch("/api/rfqs/{case_id}/lines/{line_id}", response_model=RFQLineItem)
def update_rfq_line(
    case_id: str,
    line_id: str,
    update: RFQLineItemUpdateRequest,
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    update.actor = current_user.name
    try:
        updated = case_manager.update_line(case_id, line_id, update)
    except ValueError as e:
        if "conflict" in str(e).lower() or "stale" in str(e).lower():
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(e))
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    if not updated:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Line item not found.")
    return updated


@app.post(
    "/api/rfqs/{case_id}/lines/{line_id}/rematch",
    response_model=RFQLineItem,
)
def rematch_rfq_line(
    case_id: str,
    line_id: str,
    request: RematchLineRequest,
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """Refresh one line's catalog match and require the reviewer to approve it again."""
    if case_manager.get_case(case_id) is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")
    request.actor = current_user.name
    try:
        return case_manager.rematch_line(case_id, line_id, request)
    except ValueError as exc:
        code = status.HTTP_404_NOT_FOUND if "not found" in str(exc).lower() else status.HTTP_409_CONFLICT
        raise HTTPException(status_code=code, detail=str(exc)) from exc


@app.post("/api/rfqs/{case_id}/review")
def batch_review_lines(
    case_id: str,
    request: BatchReviewRequest,
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    case = case_manager.get_case(case_id)
    if not case:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")
    request.actor = current_user.name
    try:
        count = case_manager.batch_review(case_id, request)
    except ValueError as exc:
        code = (
            status.HTTP_409_CONFLICT
            if "stale" in str(exc).lower() or "conflict" in str(exc).lower()
            else status.HTTP_400_BAD_REQUEST
        )
        raise HTTPException(status_code=code, detail=str(exc)) from exc
    return {"status": "success", "updated_count": count}


@app.get("/api/rfqs/{case_id}/review-plan", response_model=ReviewPlan)
def get_rfq_review_plan(
    case_id: str,
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    try:
        return case_manager.get_review_plan(case_id)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exc)) from exc


@app.post("/api/rfqs/{case_id}/review-plan/approve-remaining")
def approve_remaining_validated_lines(
    case_id: str,
    request: ApproveRemainingRequest,
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    case = case_manager.get_case(case_id)
    if not case:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")
    try:
        count = case_manager.approve_remaining_clean(
            case_id, current_user.name, request.expected_versions
        )
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
    return {"status": "success", "updated_count": count}


@app.get("/api/rfqs/{case_id}/audit", response_model=List[AuditEvent])
def get_rfq_audit_trail(case_id: str, current_user: PilotUser = Depends(get_current_user)):
    case = case_manager.get_case(case_id)
    if not case:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")
    return case_manager.get_audit_trail(case_id)


@app.delete("/api/rfqs/{case_id}")
def delete_rfq(
    case_id: str,
    current_user: PilotUser = Depends(require_role(["admin"])),
):
    """Deletes a case, its source data, and generated exports (admin only)."""
    if not case_manager.get_case(case_id):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")
    try:
        export_generator.delete_case_exports(case_id)
    except OSError as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Generated RFQ exports could not be deleted safely; the case was retained.",
        ) from exc
    success = case_manager.delete_case(case_id)
    if not success:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")
    return {"status": "success", "deleted_case_id": case_id}


# ---------------------------------------------------------------------------
# Export Endpoints
# ---------------------------------------------------------------------------

@app.get("/api/rfqs/{case_id}/exports", response_model=List[ExportDocumentSummary])
def list_rfq_exports(
    case_id: str,
    current_user: PilotUser = Depends(get_current_user),
):
    """List the generated worksheet files belonging to one RFQ case."""
    if not case_manager.get_case(case_id):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")
    return export_generator.list_exports(case_id)


@app.post("/api/rfqs/{case_id}/exports", response_model=ExportManifest)
def export_rfq_worksheet(
    case_id: str,
    request: ExportRequest,
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """Generates deterministic XLSX/CSV quote worksheet and email draft."""
    case = case_manager.get_case(case_id)
    if not case:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")

    lines = case_manager.get_lines(case_id)
    request.actor = current_user.name
    try:
        manifest = export_generator.generate_export(case, lines, request)
        # Record audit event
        case_manager._record_audit(
            case_id=case_id,
            actor=request.actor,
            action=AuditAction.EXPORTED,
            notes=f"Exported {manifest.format.upper()} worksheet ({manifest.file_name}) - {manifest.approved_lines} approved lines.",
        )
        return manifest
    except ExportValidationError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@app.get("/api/rfqs/{case_id}/exports/{export_id}")
def download_export(
    case_id: str,
    export_id: str,
    current_user: PilotUser = Depends(get_current_user),
):
    """Downloads the generated export artifact with strict case-scoping verification."""
    # Validate export_id format to prevent path traversal via URL path
    import re as _re
    if not _re.fullmatch(r"EXP-[A-F0-9]{8}", export_id):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Export not found.")
    export_dir = Path(export_generator.base_export_dir) / export_id
    if not export_dir.exists():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Export not found.")

    manifest_path = export_dir / "manifest.json"
    if manifest_path.exists():
        import json
        with open(manifest_path, "r", encoding="utf-8") as mf:
            manifest_data = json.load(mf)
            if manifest_data.get("case_id") != case_id:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Export '{export_id}' does not belong to case '{case_id}'.",
                )

    files = [f for f in export_dir.iterdir() if f.name != "manifest.json"]
    if not files:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Export artifact missing.")

    artifact = files[0]
    if case_id not in artifact.name:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Export '{export_id}' does not belong to case '{case_id}'.",
        )

    from fastapi.responses import FileResponse
    media_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if artifact.suffix == ".xlsx"
        else "text/csv"
    )
    return FileResponse(path=str(artifact), filename=artifact.name, media_type=media_type)


@app.get("/api/rfqs/{case_id}/email-draft")
def get_unsent_email_draft(case_id: str, current_user: PilotUser = Depends(get_current_user)):
    """Returns the unsent customer clarification email draft."""
    case = case_manager.get_case(case_id)
    if not case:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")
    lines = case_manager.get_lines(case_id)
    from rfq_assistant.exports.email_draft import generate_unsent_email_draft
    draft = generate_unsent_email_draft(case, lines)
    return {"case_id": case_id, "draft": draft}


# ---------------------------------------------------------------------------
# V1 Public Contract Endpoints (OSS Stack & Pydantic AI Pipeline)
# ---------------------------------------------------------------------------

@app.post("/v1/rfqs/extract")
async def v1_extract_rfq(
    file: Optional[UploadFile] = File(None),
    pasted_text: Optional[str] = Form(None),
    title: Optional[str] = Form(None),
    customer_name: Optional[str] = Form(None),
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """V1 Extract: Scans attachment with ClamAV, isolates untrusted data, extracts RFQ lines."""
    if not file and not pasted_text:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Must provide an attachment file or pasted text.",
        )

    if file:
        content = await _read_bounded_upload(file)
        filename = file.filename or "upload.dat"
        # 1. Attachment security scan
        scan = clamav_scanner.scan_bytes(content, filename=filename)
        if _scan_blocks_intake(scan):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    "Security scan unavailable. Processing is blocked in protected mode."
                    if not scan.scanned and clamav_scanner.protected_mode
                    else f"Security alert: Malicious attachment quarantined ({scan.threat_name}). Processing blocked."
                ),
            )

        case = case_manager.create_case_from_upload(
            filename=filename,
            content=content,
            title=title,
            customer_label=customer_name,
            actor=current_user.name,
            intake_warnings=_unscanned_warning(scan),
        )
    else:
        content = (pasted_text or "").encode("utf-8")
        scan = clamav_scanner.scan_bytes(content, filename="pasted_text.txt")
        if _scan_blocks_intake(scan):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    "Security scan unavailable. Processing is blocked in protected mode."
                    if not scan.scanned and clamav_scanner.protected_mode
                    else f"Security alert: Malicious content quarantined ({scan.threat_name})."
                ),
            )
        try:
            case = case_manager.create_case_from_text(
                raw_text=pasted_text or "",
                title=title,
                customer_label=customer_name,
                actor=current_user.name,
                intake_warnings=_unscanned_warning(scan),
            )
        except (UploadValidationError, DocumentParserError) as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc

    lines = case_manager.get_lines(case.case_id)
    return {
        "case": case,
        "lines": lines,
        "security_scan": {
            "engine": scan.engine,
            "is_clean": scan.is_clean,
            "scanned": getattr(scan, "scanned", True),
            "file_digest": scan.file_digest,
        },
    }


@app.post("/v1/rfqs/{case_id}/match")
def v1_match_rfq(
    case_id: str,
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """V1 Match: Runs exact catalog matching and RapidFuzz candidate retrieval."""
    case = case_manager.get_case(case_id)
    if not case:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")

    lines = case_manager.get_lines(case_id)
    rapidfuzz_engine.index_catalog(catalog_repo.list_all())

    matched_results = []
    for line in lines:
        candidates = rapidfuzz_engine.find_candidates(
            requested_mpn=line.requested_mpn or "",
            requested_oem=line.requested_manufacturer,
        )
        matched_results.append({
            "line_id": line.line_id,
            "line_number": line.line_number,
            "requested_mpn": line.requested_mpn,
            "current_status": line.match_status,
            "candidate_count": len(candidates),
            "candidates": [
                {
                    "internal_sku": c[0].internal_sku,
                    "mpn": c[0].manufacturer_part_number,
                    "manufacturer": c[0].manufacturer,
                    "score": c[1],
                    "reason": c[2],
                }
                for c in candidates
            ],
        })

    return {"case_id": case_id, "results": matched_results}


@app.post("/v1/rfqs/{case_id}/draft-quote")
def v1_draft_quote(
    case_id: str,
    currency: str = "USD",
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """V1 Draft Quote: Compiles ProposedQuote with provenance, authoritative pricing, and draft hash."""
    case = case_manager.get_case(case_id)
    if not case:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")

    lines = case_manager.get_lines(case_id)
    quote_lines: List[QuoteLine] = []
    buyer_quote_inputs = []
    distributor_quote_records = []
    quote_lines_by_number = {}
    accepted_updates = catalog_update_service.list_proposals()

    for idx, line in enumerate(lines, start=1):
        # Look up price and provenance from authoritative catalog if matched
        catalog_item = (
            catalog_repo.get_by_sku(line.selected_catalog_sku)
            if line.selected_catalog_sku
            else None
        )

        unit_price = (
            Decimal(str(catalog_item.list_price))
            if catalog_item and catalog_item.list_price is not None
            else None
        )
        source_price_id = (
            f"CAT-{catalog_item.catalog_version}-{catalog_item.internal_sku}"
            if catalog_item and unit_price is not None
            else None
        )

        sources = []
        if line.source_locator:
            sources.append(
                SourceRef(
                    file_id=(
                        line.source_locator.source_file_id
                        or (case.source_file.file_id if case.source_file else "unknown")
                    ),
                    page=line.source_locator.page_number,
                    sheet=line.source_locator.sheet_name,
                    text_excerpt=line.source_locator.text_snippet,
                    bbox=(list(line.source_locator.page_bbox) if line.source_locator.page_bbox else None),
                )
            )

        match_status_str = "exact" if line.match_status.value == "EXACT_MATCH" else "unmatched"

        quote_line = QuoteLine(
                line_no=idx,
                requested_mpn=line.requested_mpn or "UNKNOWN",
                requested_manufacturer=line.requested_manufacturer,
                catalogue_part_id=catalog_item.internal_sku if catalog_item else None,
                description=line.description or (catalog_item.description if catalog_item else None),
                quantity=Decimal(str(line.quantity)) if line.quantity is not None else None,
                unit=line.unit or "EA",
                unit_price=unit_price,
                currency=(catalog_item.currency if catalog_item and catalog_item.currency else currency),
                source_price_id=source_price_id,
                match_status=match_status_str,
                requires_human_review=line.review_decision.value != "APPROVED",
                sources=sources,
            )
        quote_lines.append(quote_line)
        quote_lines_by_number[idx] = quote_line

        if (
            catalog_item
            and unit_price is not None
            and line.quantity is not None
            and Decimal(str(line.quantity)) > 0
            and match_status_str == "exact"
        ):
            source_id = f"CAT-{catalog_item.catalog_version}-{catalog_item.internal_sku}"
            evidence_text = (
                f"Internal SKU {catalog_item.internal_sku}; manufacturer {catalog_item.manufacturer}; "
                f"MPN {catalog_item.manufacturer_part_number}; description {catalog_item.description}; "
                f"list price {unit_price} {catalog_item.currency or currency}; stock {catalog_item.stock_qty}"
            )
            source_locator = {}
            for proposal in accepted_updates:
                current = proposal.current_item
                if (
                    proposal.review_status == "accepted"
                    and current
                    and current.internal_sku == catalog_item.internal_sku
                    and current.catalog_version == catalog_item.catalog_version
                    and current.list_price is not None
                    and Decimal(str(current.list_price)) == unit_price
                    and proposal.evidence_quote
                ):
                    source_id = f"CATUP-{proposal.proposal_id}"
                    evidence_text = proposal.evidence_quote
                    source_locator = proposal.source_locator.model_dump(mode="json")
                    break

            buyer_source_id = (
                line.source_locator.source_file_id
                or (case.source_file.file_id if case.source_file else f"RFQ-{case_id}")
            )
            buyer_text = line.raw_text or (line.source_locator.text_snippet if line.source_locator else "") or ""
            if buyer_text:
                buyer_quote_inputs.append({
                    "line_no": idx,
                    "source_id": buyer_source_id,
                    "text": buyer_text,
                    "requested_mpn": line.requested_mpn or "UNKNOWN",
                    "quantity": str(line.quantity),
                    "locator": line.source_locator.model_dump(mode="json") if line.source_locator else {},
                })
                distributor_quote_records.append({
                    "line_no": idx,
                    "source_id": source_id,
                    "evidence_text": evidence_text,
                    "unit_price": str(unit_price),
                    "currency": catalog_item.currency or currency,
                    "locator": source_locator,
                })
                quote_line.sources.append(SourceRef(
                    file_id=source_id,
                    page=source_locator.get("page_number"),
                    sheet=source_locator.get("sheet_name"),
                    cell_range=", ".join((line.source_locator.field_source_cells or {}).values()) or None,
                    text_excerpt=evidence_text[:500],
                    bbox=source_locator.get("page_bbox"),
                ))

    quote_id = f"QUOTE-{case_id.replace('RFQ-', '')}"
    proposal_status: Literal["grounded", "not_configured", "failed", "no_retrieved_prices"] = (
        "not_configured" if distributor_quote_records else "no_retrieved_prices"
    )
    proposal_warnings = []
    extractor = case_manager.llm_extractor
    if extractor and distributor_quote_records:
        try:
            grounded_sources = GroundedQuoteProposalService(extractor.client).propose(
                buyer_lines=buyer_quote_inputs,
                distributor_records=distributor_quote_records,
            )
            for line_no, refs in grounded_sources.items():
                quote_lines_by_number[line_no].sources = refs
            proposal_status = "grounded"
        except QuoteProposalError as exc:
            proposal_status = "failed"
            proposal_warnings.append(str(exc))
    elif distributor_quote_records:
        proposal_warnings.append(
            "GPT-6 Luna quote evidence mapping is not configured. Values were copied only from the active catalog; configure Luna before treating this as a model-prepared draft."
        )

    draft = ProposedQuote(
        quote_id=quote_id,
        case_id=case_id,
        customer_name=case.customer_label or "Valued Customer",
        currency=currency,
        source_state_hash=_rfq_source_snapshot_hash(case_id),
        total_amount=calculate_quote_total(quote_lines),
        lines=quote_lines,
        quote_proposal_status=proposal_status,
        quote_proposal_warnings=proposal_warnings,
    )

    registered = commercial_service.register_draft(draft)
    return registered


@app.get("/v1/quote-drafts/{quote_id}")
def v1_get_quote_draft(quote_id: str, current_user: PilotUser = Depends(get_current_user)):
    """Retrieves a quotation draft with current hash and validation errors."""
    draft = commercial_service.get_draft(quote_id)
    if not draft:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Quote draft not found.")
    return draft


@app.get("/v1/rfqs/{case_id}/quote-draft")
def v1_get_case_quote_draft(case_id: str, current_user: PilotUser = Depends(get_current_user)):
    """Returns the saved quote draft for a case, if one has been prepared."""
    if not case_manager.get_case(case_id):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="RFQ case not found.")
    quote_id = f"QUOTE-{case_id.replace('RFQ-', '')}"
    draft = commercial_service.get_draft(quote_id)
    if not draft:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Quote draft not found.")
    return draft


@app.post("/v1/quote-drafts/{quote_id}/approve")
def v1_approve_quote_draft(
    quote_id: str,
    expected_hash: str = Form(...),
    approver_name: str = Form("Inside Sales"),
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """Version-bound human approval of a quotation draft."""
    _require_fresh_quote_source(quote_id)
    effective_approver = current_user.name
    try:
        record = commercial_service.approve_draft(
            quote_id=quote_id,
            expected_hash=expected_hash,
            approver_name=effective_approver,
        )
        return {"status": "success", "approval_record": record}
    except ApprovalGateError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@app.post("/v1/quote-drafts/{quote_id}/reject")
def v1_reject_quote_draft(
    quote_id: str,
    reason: str = Form(...),
    approver_name: str = Form("Inside Sales"),
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """Rejects quotation draft with explicit reason."""
    effective_approver = current_user.name
    try:
        record = commercial_service.reject_draft(
            quote_id=quote_id,
            approver_name=effective_approver,
            reason=reason,
        )
        return {"status": "success", "rejection_record": record}
    except ApprovalGateError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@app.post("/v1/approved-quotes/{quote_id}/send")
def v1_send_approved_quote(
    quote_id: str,
    expected_hash: str = Form(...),
    sender_name: str = Form("Inside Sales"),
    current_user: PilotUser = Depends(require_role(["admin", "reviewer"])),
):
    """Transmits quote strictly after human approval verification and hash validation."""
    _require_fresh_quote_source(quote_id)
    effective_sender = current_user.name
    try:
        result = commercial_service.send_approved_quote(
            quote_id=quote_id,
            expected_hash=expected_hash,
            sender_name=effective_sender,
        )
        return {"status": "success", "transmission": result}
    except ApprovalGateError as e:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail=str(e))

