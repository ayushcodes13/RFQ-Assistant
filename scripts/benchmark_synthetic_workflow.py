"""Measure the local deterministic RFQ workflow with generated synthetic inputs.

This is a profiling aid, not a customer-accuracy or production-capacity claim.
It performs no network/model calls and writes only to a caller-selected output.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import platform
import statistics
import subprocess
import tempfile
import time
import tracemalloc
from datetime import datetime, timezone
from io import StringIO
from pathlib import Path
from typing import Any

import pymupdf as fitz
from openpyxl import Workbook

from rfq_assistant.cases.manager import CaseManager
from rfq_assistant.exports import ExportGenerator
from rfq_assistant.ingestion.storage import DocumentStorage
from rfq_assistant.matching.catalog import CatalogRepository
from rfq_assistant.persistence import SQLiteStateStore
from rfq_assistant.schemas.enums import CaseStatus, ReviewDecision
from rfq_assistant.schemas.models import ExportRequest, RFQLineItemUpdate

FORMATS = ("csv", "xlsx", "pdf")
LINE_COUNTS = (10, 50, 100, 500)
ATTACHMENT_COUNTS = (1, 2, 5)
CATALOG_HEADERS = (
    "internal_sku",
    "manufacturer",
    "manufacturer_part_number",
    "description",
    "list_price",
    "currency",
    "stock_qty",
)
RFQ_HEADERS = ("Manufacturer", "MPN", "Qty", "Unit", "Description")
SYNTHETIC_MANUFACTURER = "Schneider Electric"


def _row(index: int) -> tuple[str, str, int, str, str]:
    mpn = f"SYN-PN-{index:06d}"
    return (
        SYNTHETIC_MANUFACTURER,
        mpn,
        1,
        "EA",
        f"Synthetic control component {index}",
    )


def _catalog_csv(line_count: int) -> bytes:
    buffer = StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(CATALOG_HEADERS)
    for index in range(1, line_count + 1):
        manufacturer, mpn, _, _, description = _row(index)
        writer.writerow(
            (
                f"SYN-SKU-{index:06d}",
                manufacturer,
                mpn,
                description,
                "1.25",
                "USD",
                1000,
            )
        )
    return buffer.getvalue().encode("utf-8")


def _csv_document(indices: range) -> bytes:
    buffer = StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(RFQ_HEADERS)
    writer.writerows(_row(index) for index in indices)
    return buffer.getvalue().encode("utf-8")


def _xlsx_document(indices: range) -> bytes:
    from io import BytesIO

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Synthetic RFQ")
    sheet.append(RFQ_HEADERS)
    for index in indices:
        sheet.append(_row(index))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_document(indices: range) -> bytes:
    pdf = fitz.open()
    page = pdf.new_page()
    y = 38

    def add_heading(target_page, top: float) -> float:
        target_page.insert_text((36, top), "SYNTHETIC RFQ - generated benchmark data only", fontsize=9)
        target_page.insert_text((36, top + 14), "Line OEM / Brand Part Number Description Qty UOM", fontsize=8)
        return top + 30

    y = add_heading(page, y)
    for index in indices:
        if y > page.rect.height - 28:
            page = pdf.new_page()
            y = add_heading(page, 38)
        manufacturer, mpn, quantity, unit, description = _row(index)
        text = f"{index} {manufacturer} {mpn} {description} {quantity} {unit}"
        page.insert_text((36, y), text, fontsize=8)
        y += 12
    content = pdf.tobytes()
    pdf.close()
    return content


def _document(filename: str, indices: range) -> bytes:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return _csv_document(indices)
    if suffix == ".xlsx":
        return _xlsx_document(indices)
    if suffix == ".pdf":
        return _pdf_document(indices)
    raise ValueError(f"Unsupported synthetic format: {suffix}")


def _percentile(values: list[float], quantile: float) -> float:
    ordered = sorted(values)
    return ordered[max(0, math.ceil(quantile * len(ordered)) - 1)]


def _summarize(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    buckets: dict[tuple[int, int, str], list[dict[str, Any]]] = {}
    for row in rows:
        key = (row["requested_lines"], row["attachment_count"], row["primary_format"])
        buckets.setdefault(key, []).append(row)
    result = []
    metrics = (
        "intake_ms",
        "queue_wait_ms",
        "parse_match_ms",
        "review_ms",
        "export_ms",
        "full_workflow_ms",
        "python_peak_bytes",
        "review_exceptions_remaining",
        "unresolved_lines_omitted",
    )
    for (lines, attachments, fmt), trials in sorted(buckets.items()):
        group: dict[str, Any] = {
            "requested_lines": lines,
            "attachment_count": attachments,
            "primary_format": fmt,
            "repetitions": len(trials),
            "successes": sum(trial["success"] for trial in trials),
            "failures": sum(not trial["success"] for trial in trials),
            "extracted_lines": sum(trial["extracted_lines"] for trial in trials),
        }
        for metric in metrics:
            values = [float(trial[metric]) for trial in trials]
            group[f"{metric}_p50"] = statistics.median(values)
            group[f"{metric}_p95"] = _percentile(values, 0.95)
        result.append(group)
    return result


def _run_trial(
    line_count: int, attachment_count: int, primary_format: str
) -> dict[str, Any]:
    with tempfile.TemporaryDirectory(prefix="rfq-synthetic-benchmark-") as temporary:
        root = Path(temporary)
        store = SQLiteStateStore(root / "state.sqlite3")
        catalog = CatalogRepository(state_store=store)
        catalog.replace_csv(_catalog_csv(line_count))
        storage = DocumentStorage(root / "sources")
        manager = CaseManager(catalog=catalog, storage=storage, state_store=store)
        exporter = ExportGenerator(catalog, root / "exports")

        # Split requested rows as evenly as possible while retaining every row.
        quotient, remainder = divmod(line_count, attachment_count)
        cursor = 1
        uploads: list[tuple[str, bytes]] = []
        for attachment_index in range(attachment_count):
            count = quotient + (1 if attachment_index < remainder else 0)
            indices = range(cursor, cursor + count)
            fmt = (
                primary_format
                if attachment_index == 0
                else FORMATS[
                    (FORMATS.index(primary_format) + attachment_index) % len(FORMATS)
                ]
            )
            name = f"synthetic-{attachment_index + 1}.{fmt}"
            uploads.append((name, _document(name, indices)))
            cursor += count

        tracemalloc.start()
        started = time.perf_counter()
        intake_started = started
        if attachment_count == 1:
            case = manager.enqueue_case_from_upload(
                *uploads[0], title="Synthetic benchmark RFQ"
            )
        else:
            case = manager.create_case_from_uploads(
                uploads, title="Synthetic benchmark RFQ", defer_processing=True
            )
        intake_ms = (time.perf_counter() - intake_started) * 1000

        job = store.claim_ingestion_job("synthetic-benchmark-worker", lease_seconds=600)
        if job is None:
            tracemalloc.stop()
            raise RuntimeError("Synthetic benchmark job could not be claimed.")
        queue_wait_ms = float(job["queue_wait_ms"])
        parse_started = time.perf_counter()
        processed = manager.process_pending_case(
            case.case_id,
            job_id=job["job_id"],
            worker_id="synthetic-benchmark-worker",
            queue_wait_ms=queue_wait_ms,
        )
        if not store.finish_ingestion_job(job["job_id"], "synthetic-benchmark-worker"):
            tracemalloc.stop()
            raise RuntimeError(
                "Synthetic benchmark worker lost its lease before completion."
            )
        parse_match_ms = (time.perf_counter() - parse_started) * 1000
        lines = manager.get_lines(case.case_id)
        exact_matches = sum(line.match_status.value == "EXACT_MATCH" for line in lines)
        success = (
            processed.status == CaseStatus.NEEDS_REVIEW
            and len(lines) == line_count
            and exact_matches == line_count
        )
        if not success:
            _, python_peak_bytes = tracemalloc.get_traced_memory()
            tracemalloc.stop()
            return {
                "requested_lines": line_count,
                "attachment_count": attachment_count,
                "primary_format": primary_format,
                "formats": [Path(name).suffix[1:] for name, _ in uploads],
                "input_bytes": sum(len(content) for _, content in uploads),
                "extracted_lines": len(lines),
                "exact_matches": exact_matches,
                "review_sample_lines": 0,
                "review_exceptions_remaining": 0,
                "export_lines": 0,
                "unresolved_lines_omitted": line_count,
                "export_is_draft": False,
                "success": False,
                "failure_reason": f"status={processed.status.value}, extracted={len(lines)}, exact_matches={exact_matches}",
                "intake_ms": intake_ms,
                "queue_wait_ms": queue_wait_ms,
                "parse_match_ms": parse_match_ms,
                "review_ms": 0.0,
                "export_ms": 0.0,
                "full_workflow_ms": (time.perf_counter() - started) * 1000,
                "python_peak_bytes": python_peak_bytes,
            }
        review_started = time.perf_counter()
        plan = manager.get_review_plan(case.case_id)
        by_id = {line.line_id: line for line in lines}
        for line_id in plan.sample_line_ids:
            line = by_id[line_id]
            manager.update_line(
                case.case_id,
                line_id,
                RFQLineItemUpdate(
                    expected_version=line.version,
                    review_decision=ReviewDecision.APPROVED,
                    actor="Synthetic benchmark reviewer",
                ),
            )
        refreshed = manager.get_lines(case.case_id)
        if manager.get_review_plan(case.case_id).remaining_line_ids:
            manager.approve_remaining_clean(
                case.case_id,
                "Synthetic benchmark reviewer",
                {line.line_id: line.version for line in refreshed},
            )
        review_ms = (time.perf_counter() - review_started) * 1000

        export_started = time.perf_counter()
        export = exporter.generate_export(
            processed,
            manager.get_lines(case.case_id),
            ExportRequest(
                case_id=case.case_id,
                actor="Synthetic benchmark reviewer",
                format="csv",
                allow_partial=True,
            ),
        )
        export_ms = (time.perf_counter() - export_started) * 1000
        _, python_peak_bytes = tracemalloc.get_traced_memory()
        tracemalloc.stop()
        full_workflow_ms = (time.perf_counter() - started) * 1000

        source_bytes = sum(len(content) for _, content in uploads)
        return {
            "requested_lines": line_count,
            "attachment_count": attachment_count,
            "primary_format": primary_format,
            "formats": [Path(name).suffix[1:] for name, _ in uploads],
            "input_bytes": source_bytes,
            "extracted_lines": len(lines),
            "exact_matches": exact_matches,
            "review_sample_lines": len(plan.sample_line_ids),
            "review_exceptions_remaining": len(plan.exception_line_ids),
            "export_lines": export.total_lines,
            "unresolved_lines_omitted": export.unresolved_lines_omitted,
            "export_is_draft": export.is_draft,
            "success": success and export.total_lines == line_count and export.is_draft,
            "failure_reason": None
            if success
            else f"status={processed.status.value}, extracted={len(lines)}",
            "intake_ms": intake_ms,
            "queue_wait_ms": queue_wait_ms,
            "parse_match_ms": parse_match_ms,
            "review_ms": review_ms,
            "export_ms": export_ms,
            "full_workflow_ms": full_workflow_ms,
            "python_peak_bytes": python_peak_bytes,
        }


def run_benchmark(repetitions: int) -> dict[str, Any]:
    trials = []
    for line_count in LINE_COUNTS:
        for attachment_count in ATTACHMENT_COUNTS:
            for primary_format in FORMATS:
                for repetition in range(repetitions):
                    row = _run_trial(line_count, attachment_count, primary_format)
                    row["repetition"] = repetition + 1
                    trials.append(row)
    elapsed = sum(row["full_workflow_ms"] for row in trials) / 1000
    total_lines = sum(row["requested_lines"] for row in trials)
    return {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "python_version": platform.python_version(),
        "system": platform.platform(),
        "source_revision": subprocess.check_output(
            ["git", "rev-parse", "HEAD"], text=True
        ).strip(),
        "scope": "Synthetic local deterministic path; no Luna or OCR calls; no network; single process.",
        "measurement_notes": [
            "Three repetitions per cell make p95 a nearest-rank estimate equal to the maximum; treat it as descriptive, not a stable tail-latency SLO.",
            "Peak memory is Python allocation peak from tracemalloc; it excludes native allocations and total process RSS.",
            "Input files are generated before tracemalloc starts, so the peak excludes pre-existing request buffers and service/database/catalog startup allocations.",
            "PDF cases are generated native-text PDFs. Scanned PDF/OCR and Luna latency/token use are not measured.",
            "Each trial uses a fresh temporary SQLite database, local source storage, catalog, and export directory.",
            "Review timing measures the plan and synthetic approval services, not human reading or correction time; remaining exceptions are recorded and the draft export explicitly allows partial output.",
        ],
        "matrix": {
            "line_counts": LINE_COUNTS,
            "attachment_counts": ATTACHMENT_COUNTS,
            "primary_formats": FORMATS,
            "repetitions": repetitions,
            "trial_count": len(trials),
        },
        "aggregate": {
            "successes": sum(row["success"] for row in trials),
            "failures": sum(not row["success"] for row in trials),
            "failure_rate": sum(not row["success"] for row in trials) / len(trials),
            "throughput_lines_per_second": total_lines / elapsed if elapsed else 0,
        },
        "summary": _summarize(trials),
        "trials": trials,
    }


def render_markdown(report: dict[str, Any]) -> str:
    matrix = report["matrix"]
    aggregate = report["aggregate"]
    lines = [
        "# Synthetic End-to-End Workflow Performance",
        "",
        f"Generated: `{report['generated_at_utc']}`",
        "",
        f"Python: `{report['python_version']}`",
        "",
        f"System: `{report['system']}`",
        "",
        f"Benchmark code revision: `{report['source_revision']}`",
        "",
        f"Scope: {report['scope']}",
        "",
        f"Matrix: {matrix['trial_count']} trials; line counts {matrix['line_counts']}; attachment counts {matrix['attachment_counts']}; primary formats {', '.join(matrix['primary_formats'])}; repetitions per cell {matrix['repetitions']}.",
        "",
        f"Overall: {aggregate['successes']} successful, {aggregate['failures']} failed, failure rate `{aggregate['failure_rate']:.3f}`, throughput `{aggregate['throughput_lines_per_second']:.1f}` requested lines/second (sum of serial trial durations).",
        "",
        "| Lines | Attachments | Primary format | Intake p50/p95 ms | Queue p50/p95 ms | Parse+match p50/p95 ms | Review p50/p95 ms | Export p50/p95 ms | Full workflow p50/p95 ms | Python peak p50/p95 MiB | Review exceptions p50/p95 | Export omissions p50/p95 | Success/fail |",
        "|---:|---:|:---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for row in report["summary"]:

        def pair(
            metric: str,
            scale: float = 1.0,
            digits: int = 1,
            row: dict[str, Any] = row,
        ) -> str:
            return f"{row[metric + '_p50'] / scale:.{digits}f}/{row[metric + '_p95'] / scale:.{digits}f}"

        lines.append(
            f"| {row['requested_lines']} | {row['attachment_count']} | {row['primary_format']} | "
            f"{pair('intake_ms')} | {pair('queue_wait_ms')} | {pair('parse_match_ms')} | "
            f"{pair('review_ms')} | {pair('export_ms')} | {pair('full_workflow_ms')} | "
            f"{pair('python_peak_bytes', 1024 * 1024, 2)} | {pair('review_exceptions_remaining', 1, 0)} | "
            f"{pair('unresolved_lines_omitted', 1, 0)} | {row['successes']}/{row['failures']} |"
        )
    lines.extend(["", "## Limits", ""])
    lines.extend(f"- {note}" for note in report["measurement_notes"])
    lines.extend(
        [
            "- Results describe this local synthetic run only. They do not establish customer accuracy, hosted capacity, or production SLOs.",
            "- No optimization is justified by this run alone; inspect the phase distributions and repeat on the intended deployment hardware before selecting a bottleneck.",
            "",
        ]
    )
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repetitions", type=int, default=3)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("tests/end_to_end_audit/results/workflow_performance_latest.json"),
    )
    args = parser.parse_args()
    if args.repetitions < 1 or args.repetitions > 20:
        parser.error("--repetitions must be from 1 through 20")
    report = run_benchmark(args.repetitions)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    markdown_path = args.output.with_suffix(".md")
    markdown_path.write_text(render_markdown(report), encoding="utf-8")
    print(f"Wrote synthetic benchmark JSON: {args.output}")
    print(f"Wrote synthetic benchmark report: {markdown_path}")
    print(json.dumps(report["aggregate"], sort_keys=True))


if __name__ == "__main__":
    main()
